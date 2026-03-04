#!/usr/bin/env python3
"""Bootstrap a localization workbook from scoped GameMaker text sources."""

from __future__ import annotations

import argparse
import re
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Optional

from openpyxl import Workbook

SHEET_NAME = "strings"
SCHEMA_LEGACY = "legacy"
SCHEMA_MINIMAL = "minimal"

LEGACY_COLUMNS = [
    "key",
    "group",
    "context",
    "source_file",
    "source_line",
    "source_ref",
    "english_en",
    "korean_ko",
    "placeholders",
    "status",
    "notes",
    "active",
    "char_limit",
]

MINIMAL_COLUMNS = [
    "ref",
    "en",
    "ko",
]

STRING_RE = re.compile(r'"((?:[^"\\]|\\.)*)"')
LOC_CALL_RE = re.compile(r'Loc_T\(\s*"((?:[^"\\]|\\.)*)"\s*,\s*"((?:[^"\\]|\\.)*)"')
PLACEHOLDER_RE = re.compile(r"\{([A-Za-z0-9_]+)\}")

INTERNAL_TOKENS = {
    "menu_up",
    "menu_down",
    "menu_left",
    "menu_right",
    "confirm",
    "interact",
    "cancel",
    "ui_move",
    "ui_confirm",
    "ui_back",
    "ui_openclose",
    "ui_open",
    "ui_close",
    "linear",
    "smoothstep",
    "default",
    "main",
    "cutscene",
    "difficulty",
    "settings",
    "to_main",
    "start_intro",
    "load",
    "save",
    "bed",
    "open_save",
    "load_slot",
    "save_and_reload_slot",
    "show_no_saves_message",
    "delete",
    "overwrite",
    "saved",
    "message",
    "title",
    "title_room_change",
    "gameplay_load",
    "intro",
    "ending",
    "game_over",
    "normal",
    "slow",
    "Background",
    "Instances",
    "transition",
    "none",
}


@dataclass
class Row:
    key: str
    group: str
    context: str
    source_file: str
    source_line: int
    source_ref: str
    english_en: str
    korean_ko: str = ""
    placeholders: str = ""
    status: str = "new"
    notes: str = ""
    active: int = 1
    char_limit: str = ""

    def to_sheet_row(self, schema: str) -> list:
        if schema == SCHEMA_MINIMAL:
            return [
                self.key,         # ref
                self.english_en,  # en
                self.korean_ko,   # ko
            ]

        return [
            self.key,
            self.group,
            self.context,
            self.source_file,
            self.source_line,
            self.source_ref,
            self.english_en,
            self.korean_ko,
            self.placeholders,
            self.status,
            self.notes,
            self.active,
            self.char_limit,
        ]


class Catalog:
    def __init__(self) -> None:
        self.rows: "OrderedDict[str, Row]" = OrderedDict()

    def add(self, row: Row) -> None:
        if row.key in self.rows:
            existing = self.rows[row.key]
            if existing.english_en != row.english_en:
                existing.notes = (existing.notes + " | " if existing.notes else "") + (
                    f"Duplicate key with conflicting English at {row.source_file}:{row.source_line}"
                )
            return
        self.rows[row.key] = row

    def sorted_rows(self) -> list[Row]:
        return [self.rows[k] for k in sorted(self.rows.keys())]


def read_lines(path: Path) -> list[str]:
    return path.read_text(encoding="utf-8").splitlines()


def rel(root: Path, path: Path) -> str:
    return path.relative_to(root).as_posix()


def decode_gml_string(raw: str) -> str:
    return (
        raw.replace(r"\\", "\\")
        .replace(r"\"", '"')
        .replace(r"\n", "\n")
        .replace(r"\r", "\r")
        .replace(r"\t", "\t")
    )


def iter_literals(line: str) -> Iterable[str]:
    for m in STRING_RE.finditer(line):
        yield decode_gml_string(m.group(1))


def placeholders_for(text: str) -> str:
    seen: list[str] = []
    for m in PLACEHOLDER_RE.finditer(text):
        p = "{" + m.group(1) + "}"
        if p not in seen:
            seen.append(p)
    return ",".join(seen)


def slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")
    return s or "text"


def is_probably_player_text(value: str) -> bool:
    if not value:
        return False
    if value in INTERNAL_TOKENS:
        return False
    if value.startswith(("ui_", "menu_", "skillbook_", "enemy_skill_", "dialogue_")):
        return False
    if re.fullmatch(r"[a-z0-9_]+", value):
        return False
    if not re.search(r"[A-Za-z]", value):
        return False
    return True


def add_row(
    catalog: Catalog,
    *,
    key: str,
    group: str,
    context: str,
    source_file: str,
    source_line: int,
    source_ref: str,
    english: str,
    active: int = 1,
) -> None:
    english = english.strip()
    if english == "":
        return
    catalog.add(
        Row(
            key=key,
            group=group,
            context=context,
            source_file=source_file,
            source_line=source_line,
            source_ref=source_ref,
            english_en=english,
            placeholders=placeholders_for(english),
            active=active,
        )
    )


def extract_dialogue_db(catalog: Catalog, root: Path) -> None:
    path = root / "scripts/scr_dialogue/scr_dialogue.gml"
    lines = read_lines(path)
    rp = rel(root, path)

    i = 0
    while i < len(lines):
        m = re.search(r'global\.dialogue_db\[\?\s*"([^"]+)"\]\s*=', lines[i])
        if not m:
            i += 1
            continue

        dialogue_id = m.group(1)
        line_index = 0
        first_line = True
        j = i
        while j < len(lines):
            literals = list(iter_literals(lines[j]))
            for lit_idx, text in enumerate(literals):
                if first_line and lit_idx == 0 and text == dialogue_id:
                    continue
                key = f"dialogue.{dialogue_id}.{line_index}"
                add_row(
                    catalog,
                    key=key,
                    group="dialogue",
                    context=f"Dialogue line {line_index} for {dialogue_id}",
                    source_file=rp,
                    source_line=j + 1,
                    source_ref=f'global.dialogue_db[? "{dialogue_id}"]',
                    english=text,
                )
                line_index += 1
            if ";" in lines[j]:
                break
            first_line = False
            j += 1

        i = j + 1


def extract_cutscenes_and_title(catalog: Catalog, root: Path) -> None:
    path = root / "objects/obj_start_controller/Create_0.gml"
    lines = read_lines(path)
    rp = rel(root, path)

    # Title menu arrays
    for idx, text in enumerate(iter_literals(lines[1])):
        key = f"menu.main.option.{idx}"
        add_row(
            catalog,
            key=key,
            group="menu",
            context="Title menu option",
            source_file=rp,
            source_line=2,
            source_ref=f"main_options[{idx}]",
            english=text,
        )

    for idx, text in enumerate(iter_literals(lines[2])):
        key = f"settings.difficulty.option.{idx}"
        add_row(
            catalog,
            key=key,
            group="settings",
            context="Difficulty option",
            source_file=rp,
            source_line=3,
            source_ref=f"difficulty_options[{idx}]",
            english=text,
        )

    current_cutscene: Optional[str] = None
    segment_index = -1
    in_lines = False
    line_index = 0

    for lineno, line in enumerate(lines, start=1):
        m_start = re.search(r"cutscene_definitions\.([A-Za-z0-9_]+)\s*=\s*\[", line)
        if m_start:
            current_cutscene = m_start.group(1)
            segment_index = -1
            in_lines = False
            continue

        if current_cutscene and re.match(r"\s*\];\s*$", line):
            current_cutscene = None
            segment_index = -1
            in_lines = False
            continue

        if not current_cutscene:
            continue

        if re.search(r"\blines\s*:\s*\[", line):
            segment_index += 1
            in_lines = True
            line_index = 0
            continue

        if in_lines:
            for text in iter_literals(line):
                key = f"cutscene.{current_cutscene}.{segment_index}.{line_index}"
                add_row(
                    catalog,
                    key=key,
                    group="cutscene",
                    context=f"Cutscene {current_cutscene} segment {segment_index} line {line_index}",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"cutscene_definitions.{current_cutscene}[{segment_index}].lines",
                    english=text,
                )
                line_index += 1
            if "]" in line:
                in_lines = False


def extract_item_db(catalog: Catalog, root: Path) -> None:
    path = root / "scripts/scr_db_item/scr_db_item.gml"
    lines = read_lines(path)
    rp = rel(root, path)

    current_item_id: Optional[str] = None
    in_item_block = False

    for lineno, line in enumerate(lines, start=1):
        id_match = re.search(r"\bid:\s*([A-Za-z0-9_]+)\s*,", line)
        if id_match:
            current_item_id = id_match.group(1)
            in_item_block = True

        name_match = re.search(r'\bname:\s*"((?:[^"\\]|\\.)*)"', line)
        if in_item_block and current_item_id and name_match:
            text = decode_gml_string(name_match.group(1))
            key = f"item.name.{slug(current_item_id)}"
            add_row(
                catalog,
                key=key,
                group="item",
                context=f"Item display name ({current_item_id})",
                source_file=rp,
                source_line=lineno,
                source_ref=f"ItemDB_Init item {current_item_id}",
                english=text,
            )

        if in_item_block and line.strip().startswith("};"):
            in_item_block = False
            current_item_id = None

        passive_match = re.search(r"case\s+(\d+)\s*:\s*return\s*\{", line)
        if passive_match and "desc:" in line:
            item_id = passive_match.group(1)
            lits = list(iter_literals(line))
            if len(lits) >= 3:
                passive_name = lits[-2]
                passive_desc = lits[-1]
                add_row(
                    catalog,
                    key=f"item.passive.name.{item_id}",
                    group="item",
                    context=f"Equip passive name for item {item_id}",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"ItemDB_GetPassiveTemplate case {item_id}",
                    english=passive_name,
                )
                add_row(
                    catalog,
                    key=f"item.passive.desc.{item_id}",
                    group="item",
                    context=f"Equip passive description for item {item_id}",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"ItemDB_GetPassiveTemplate case {item_id}",
                    english=passive_desc,
                )


def extract_skill_db(catalog: Catalog, root: Path) -> None:
    path = root / "scripts/scr_db_skill/scr_db_skill.gml"
    lines = read_lines(path)
    rp = rel(root, path)

    current_skill_id: Optional[str] = None
    in_skill_block = False

    for lineno, line in enumerate(lines, start=1):
        map_match = re.search(r"global\.skill_db\[\?\s*([A-Za-z0-9_]+)\s*\]\s*=\s*\{", line)
        if map_match:
            current_skill_id = map_match.group(1)
            in_skill_block = True

        if in_skill_block and current_skill_id:
            name_match = re.search(r'\bname:\s*"((?:[^"\\]|\\.)*)"', line)
            if name_match:
                text = decode_gml_string(name_match.group(1))
                key = f"skill.name.{slug(current_skill_id)}"
                add_row(
                    catalog,
                    key=key,
                    group="skill",
                    context=f"Skill display name ({current_skill_id})",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"SkillDB_Init skill {current_skill_id}",
                    english=text,
                )

            use_msg_match = re.search(r'\buse_msg:\s*"((?:[^"\\]|\\.)*)"', line)
            if use_msg_match:
                text = decode_gml_string(use_msg_match.group(1))
                key = f"combat.msg.skill.use.{slug(current_skill_id)}"
                add_row(
                    catalog,
                    key=key,
                    group="combat.msg",
                    context=f"Skill use message ({current_skill_id})",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"SkillDB_Init skill {current_skill_id} use_msg",
                    english=text,
                )

        if in_skill_block and line.strip().startswith("};"):
            in_skill_block = False
            current_skill_id = None


def extract_enemy_db(catalog: Catalog, root: Path) -> None:
    path = root / "scripts/scr_db_enemy/scr_db_enemy.gml"
    lines = read_lines(path)
    rp = rel(root, path)

    current_enemy_id: Optional[str] = None
    in_enemy_block = False

    for lineno, line in enumerate(lines, start=1):
        map_match = re.search(r"global\.enemy_db\[\?\s*([A-Za-z0-9_]+)\s*\]\s*=\s*\{", line)
        if map_match:
            current_enemy_id = map_match.group(1)
            in_enemy_block = True

        if in_enemy_block and current_enemy_id:
            name_match = re.search(r'\bname:\s*"((?:[^"\\]|\\.)*)"', line)
            if name_match:
                text = decode_gml_string(name_match.group(1))
                key = f"enemy.name.{slug(current_enemy_id)}"
                add_row(
                    catalog,
                    key=key,
                    group="enemy",
                    context=f"Enemy display name ({current_enemy_id})",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"EnemyDB_Init enemy {current_enemy_id}",
                    english=text,
                )

        if in_enemy_block and line.strip().startswith("};"):
            in_enemy_block = False
            current_enemy_id = None


def extract_status_db(catalog: Catalog, root: Path) -> None:
    path = root / "scripts/scr_status/scr_status.gml"
    lines = read_lines(path)
    rp = rel(root, path)

    current_status_id: Optional[str] = None
    in_status_block = False

    for lineno, line in enumerate(lines, start=1):
        map_match = re.search(r"global\.status_db\[\?\s*([A-Za-z0-9_]+)\s*\]\s*=\s*\{", line)
        if map_match:
            current_status_id = map_match.group(1)
            in_status_block = True

        if in_status_block and current_status_id:
            name_match = re.search(r'\bname:\s*"((?:[^"\\]|\\.)*)"', line)
            if name_match:
                text = decode_gml_string(name_match.group(1))
                key = f"status.name.{slug(current_status_id)}"
                add_row(
                    catalog,
                    key=key,
                    group="status",
                    context=f"Status display name ({current_status_id})",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"StatusDB_Init status {current_status_id}",
                    english=text,
                )

        if in_status_block and line.strip().startswith("};"):
            in_status_block = False
            current_status_id = None


def extract_class_db(catalog: Catalog, root: Path) -> None:
    path = root / "scripts/scr_db_player_class/scr_db_player_class.gml"
    lines = read_lines(path)
    rp = rel(root, path)

    current_class: Optional[str] = None

    for lineno, line in enumerate(lines, start=1):
        case_match = re.search(r"case\s+([A-Za-z0-9_]+)\s*:", line)
        if case_match:
            current_class = case_match.group(1)

        name_match = re.search(r'name\s*:\s*"((?:[^"\\]|\\.)*)"', line)
        if current_class and name_match:
            text = decode_gml_string(name_match.group(1))
            key = f"class.name.{slug(current_class)}"
            add_row(
                catalog,
                key=key,
                group="class",
                context=f"Player class name ({current_class})",
                source_file=rp,
                source_line=lineno,
                source_ref=f"DB_PlayerClass case {current_class}",
                english=text,
            )


def derive_group(file_rel: str, func_name: str) -> str:
    if file_rel.endswith("scripts/scr_menu/scr_menu.gml"):
        if func_name.startswith("SettingsPopup"):
            return "settings"
        if func_name.startswith("PauseMenu"):
            return "pause"
        if func_name.startswith("ClassSelect"):
            return "class"
        return "menu"
    if file_rel.endswith("scripts/scr_save_menu/scr_save_menu.gml"):
        if func_name.startswith("BedMenu"):
            return "bed"
        return "save"
    if file_rel.endswith("scripts/scr_combat/scr_combat.gml"):
        return "combat.msg"
    if file_rel.endswith("scripts/scr_interact/scr_interact.gml"):
        return "interact"
    if file_rel.endswith("objects/obj_battle_controller/Create_0.gml"):
        return "battle.action"
    if file_rel.endswith("objects/obj_battle_controller/Step_0.gml"):
        return "battle.menu"
    if file_rel.endswith("objects/obj_battle_controller/Draw_64.gml"):
        return "battle.menu"
    if file_rel.endswith("objects/obj_start_controller/Draw_64.gml"):
        return "menu"
    if file_rel.endswith("objects/obj_start_controller/Step_0.gml"):
        return "menu"
    return "system.msg"


def include_line_context(line: str) -> bool:
    include_tokens = (
        "draw_text(",
        "Battle_Message(",
        "Menu_InvPopupOpenMessage(",
        "SaveMenu_OpenMessage(",
        "message_text",
        "options:",
        "tabs:",
        "choices:",
        "title =",
        "label =",
        "value =",
        "confirm_text",
        "cancel_text",
        "yes_label",
        "no_label",
        "ok_label",
        ".msg",
        "popup_msg",
        "exp_label",
        "array_push(lines",
        "array_push(shown",
        "array_push(out",
        "array_push(pages",
        "main_options",
        "difficulty_options",
        "battle_actions =",
        "interact_name",
        "grave_lines",
    )
    return any(token in line for token in include_tokens)


def extract_generic_ui_strings(catalog: Catalog, root: Path) -> None:
    targets = [
        root / "scripts/scr_menu/scr_menu.gml",
        root / "scripts/scr_save_menu/scr_save_menu.gml",
        root / "scripts/scr_combat/scr_combat.gml",
        root / "scripts/scr_interact/scr_interact.gml",
        root / "objects/obj_start_controller/Draw_64.gml",
        root / "objects/obj_start_controller/Step_0.gml",
        root / "objects/obj_battle_controller/Create_0.gml",
        root / "objects/obj_battle_controller/Step_0.gml",
        root / "objects/obj_battle_controller/Draw_64.gml",
        root / "objects/obj_grave/Create_0.gml",
    ]

    for path in targets:
        lines = read_lines(path)
        rp = rel(root, path)
        func_name = path.stem

        for lineno, line in enumerate(lines, start=1):
            fn_match = re.match(r"\s*function\s+([A-Za-z0-9_]+)\s*\(", line)
            if fn_match:
                func_name = fn_match.group(1)

            if not include_line_context(line):
                continue

            literals = list(iter_literals(line))
            if not literals:
                continue

            group = derive_group(rp, func_name)
            for idx, text in enumerate(literals):
                if not is_probably_player_text(text):
                    continue
                key = f"{group}.{slug(func_name)}.{lineno}.{idx}"
                add_row(
                    catalog,
                    key=key,
                    group=group,
                    context=f"{func_name} UI/message text",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"{func_name}:{lineno}",
                    english=text,
                )


def extract_interactable_create_strings(catalog: Catalog, root: Path) -> None:
    create_files = sorted((root / "objects").glob("*/Create_0.gml"))

    for path in create_files:
        lines = read_lines(path)
        rp = rel(root, path)
        obj_name = path.parent.name

        in_grave_lines = False
        grave_idx = 0

        for lineno, line in enumerate(lines, start=1):
            iname_match = re.search(r'interact_name\s*=\s*"((?:[^"\\]|\\.)*)"', line)
            if iname_match:
                text = decode_gml_string(iname_match.group(1))
                key = f"interact.name.{slug(obj_name)}"
                add_row(
                    catalog,
                    key=key,
                    group="interact",
                    context=f"Interactable display name ({obj_name})",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"{obj_name}.interact_name",
                    english=text,
                )

            if re.search(r"\bgrave_lines\s*=\s*\[", line):
                in_grave_lines = True
                grave_idx = 0
                continue

            if in_grave_lines:
                for text in iter_literals(line):
                    key = f"interact.grave.line.{slug(obj_name)}.{grave_idx}"
                    add_row(
                        catalog,
                        key=key,
                        group="interact",
                        context=f"Grave flavor line ({obj_name})",
                        source_file=rp,
                        source_line=lineno,
                        source_ref=f"{obj_name}.grave_lines[{grave_idx}]",
                        english=text,
                    )
                    grave_idx += 1

                if "]" in line:
                    in_grave_lines = False


def extract_loc_t_calls(catalog: Catalog, root: Path) -> None:
    scoped_files = [
        root / "scripts/scr_player/scr_player.gml",
        root / "scripts/scr_menu/scr_menu.gml",
        root / "scripts/scr_save_menu/scr_save_menu.gml",
        root / "scripts/scr_combat/scr_combat.gml",
        root / "scripts/scr_dialogue/scr_dialogue.gml",
        root / "scripts/scr_db_item/scr_db_item.gml",
        root / "scripts/scr_db_skill/scr_db_skill.gml",
        root / "scripts/scr_db_enemy/scr_db_enemy.gml",
        root / "scripts/scr_status/scr_status.gml",
        root / "scripts/scr_db_player_class/scr_db_player_class.gml",
        root / "scripts/scr_interact/scr_interact.gml",
        root / "objects/obj_start_controller/Create_0.gml",
        root / "objects/obj_start_controller/Step_0.gml",
        root / "objects/obj_start_controller/Draw_64.gml",
        root / "objects/obj_battle_controller/Create_0.gml",
        root / "objects/obj_battle_controller/Step_0.gml",
        root / "objects/obj_battle_controller/Draw_64.gml",
        root / "objects/obj_grave/Create_0.gml",
    ]

    for path in scoped_files:
        if not path.exists():
            continue
        lines = read_lines(path)
        rp = rel(root, path)
        func_name = path.stem

        for lineno, line in enumerate(lines, start=1):
            fn_match = re.match(r"\s*function\s+([A-Za-z0-9_]+)\s*\(", line)
            if fn_match:
                func_name = fn_match.group(1)

            for m in LOC_CALL_RE.finditer(line):
                key = decode_gml_string(m.group(1))
                fallback = decode_gml_string(m.group(2))
                group = key.split(".")[0] if "." in key else "system"
                add_row(
                    catalog,
                    key=key,
                    group=group,
                    context=f"Loc_T fallback in {func_name}",
                    source_file=rp,
                    source_line=lineno,
                    source_ref=f"{func_name}:{lineno}",
                    english=fallback,
                )


def write_workbook(catalog: Catalog, output_path: Path, schema: str = SCHEMA_MINIMAL) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    columns = MINIMAL_COLUMNS if schema == SCHEMA_MINIMAL else LEGACY_COLUMNS
    ws.append(columns)

    for row in catalog.sorted_rows():
        ws.append(row.to_sheet_row(schema))

    ws.freeze_panes = "A2"
    last_col_letter = chr(ord("A") + len(columns) - 1)
    ws.auto_filter.ref = f"A1:{last_col_letter}{ws.max_row}"

    if schema == SCHEMA_MINIMAL:
        widths = {
            "A": 52,
            "B": 72,
            "C": 72,
        }
    else:
        widths = {
            "A": 42,
            "B": 14,
            "C": 42,
            "D": 48,
            "E": 10,
            "F": 42,
            "G": 56,
            "H": 56,
            "I": 24,
            "J": 12,
            "K": 28,
            "L": 8,
            "M": 10,
        }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    wb.save(output_path)


def build_catalog(root: Path) -> Catalog:
    catalog = Catalog()
    extract_dialogue_db(catalog, root)
    extract_cutscenes_and_title(catalog, root)
    extract_item_db(catalog, root)
    extract_skill_db(catalog, root)
    extract_enemy_db(catalog, root)
    extract_status_db(catalog, root)
    extract_class_db(catalog, root)
    extract_interactable_create_strings(catalog, root)
    extract_generic_ui_strings(catalog, root)
    extract_loc_t_calls(catalog, root)
    return catalog


def main() -> None:
    parser = argparse.ArgumentParser(description="Bootstrap localization workbook from GML sources.")
    parser.add_argument("--xlsx", required=True, help="Output workbook path")
    parser.add_argument(
        "--schema",
        choices=[SCHEMA_MINIMAL, SCHEMA_LEGACY],
        default=SCHEMA_MINIMAL,
        help="Workbook schema to emit (default: minimal).",
    )
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[2]
    output_path = (root / args.xlsx).resolve() if not Path(args.xlsx).is_absolute() else Path(args.xlsx)

    catalog = build_catalog(root)
    write_workbook(catalog, output_path, args.schema)
    print(f"Wrote {len(catalog.rows)} rows to {output_path} (schema={args.schema})")


if __name__ == "__main__":
    main()
