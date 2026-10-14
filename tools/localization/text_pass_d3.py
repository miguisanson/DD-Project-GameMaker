#!/usr/bin/env python3
"""D3 text pass over datafiles/localization/game_text.xlsx (re-runnable, idempotent).

What it does, in order:
  1. Removes the leftover line-number rows made by the first bootstrap (refs like menu.draw_64.91.0 whose English is
     just another key name). No code reads them; the real keys have their own rows.
  2. Rewrites the ten grave lines (interact.grave.line.obj_grave.0-9) - the sheet had them on odd rows with key text between.
  3. Applies EDITS (English and Korean fixes), ADDS (new rows) and the global term normalisation below.
  4. Saves the workbook. Rebuild the JSON afterwards:
       python tools/localization/build_localization.py --xlsx datafiles/localization/game_text.xlsx --json datafiles/localization/localization.json

Merge rule: if the workbook conflicts on a merge, take the other branch's workbook and run this script again.
It only touches the rows named here, so rows other packets added are left alone (apart from the global term fixes).

Style decisions are in .planning/research/ko-glossary.md.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "datafiles/localization/game_text.xlsx"

JUNK_STEMS = {
    "battle.action.create_0", "battle.menu.draw_64", "battle.menu.step_0", "bed.bedmenu_draw", "class.classselect_draw",
    "combat.msg.battle_buildvictorydialoguelines", "combat.msg.battle_checkend", "combat.msg.battle_enemyact",
    "combat.msg.battle_playerattackresolvetimed", "combat.msg.battle_playerskill", "combat.msg.battle_playerstunskip",
    "combat.msg.battle_runattempt", "menu.draw_64", "menu.menu_draw", "menu.menu_handleinput", "menu.menu_useinventoryitem",
    "menu.tooltip_builditemlines", "menu.tooltip_skilleffectlines", "save.savemenu_draw", "save.savemenu_handle",
    "settings.settingspopup_draw", "system.msg.create_0",
}
JUNK_RE = re.compile(r"^(?P<stem>.+)\.\d+\.\d+$")

# "ref || en || ko" - "=" keeps the current value. Blank lines and # lines are ignored.
EDITS = """
# ---------------- battle ----------------
common.ok || = || 확인
class.line.mage || A spell caster. Great magic and MP, five skills to find. || 주문 시전자. 뛰어난 마법과 MP, 찾아야 할 스킬 5개.
item.sent.31 || = || 적에게 독 {n}중첩을 겁니다. 턴을 소모합니다.
battle.action.skill || = || 스킬
battle.action.timing.good || = || 좋음
battle.action.timing.ok || = || 무난
battle.action.timing.perfect || = || 완벽
battle.action.timing.weak || = || 약함
battle.menu.back || = || 뒤로
battle.menu.no_items || = || 사용할 수 있는 아이템이 없습니다.
battle.menu.skill_row || {name} ({mp} MP) || {name} ({mp} MP)
battle.item.other || = || 쓸모 있는 아이템입니다.
battle.hud.break || = || 브레이크!
battle.det.hits || = || 타수
battle.chip.steadfast.line || = || 기절하지 않습니다.
# ---------------- bed ----------------
bed.option.rest || = || 휴식
bed.option.heal || = || 쉬기
bed.rest_moment || = || 잠시 쉬었다.
# ---------------- class ----------------
class.name.archer || = || 궁수
class.name.class_archer || = || 궁수
class.name.class_knight || = || 기사
class.name.class_mage || = || 마법사
class.name.class_nobody || = || 무명
class.name.mage || = || 마법사
class.name.unknown || = || 알 수 없음
class.name.warrior || = || 전사
class.select.title || Select class || 직업 선택
class.lbl.five || = || 찾을 스킬 5개
class.select.become || = || {cls}의 길을 걸을까요?
# ---------------- combat log ----------------
combat.msg.applied_prefix || Applied  || 부여:
combat.msg.auto_stats || = || 자동 {summary}
combat.msg.victory_prefix || You defeated {enemy}. || {enemy}을(를) 쓰러뜨렸습니다.
combat.msg.loot_prefix || Loot:  || 전리품:
combat.msg.floor1_reaction.0 || How was that even alive?! I nearly died. || 저게 어떻게 살아 있지?! 죽을 뻔했어.
combat.msg.floor1_reaction.1 || = || 맙소사... 어서 여기서 나가야 해.
combat.msg.enemy_slain || = || {enemy}이(가) 쓰러졌습니다.
combat.msg.enemy_stunned || = || {enemy}이(가) 기절했습니다!
combat.msg.enemy_miss || = || {enemy}의 공격이 빗나갔습니다!
combat.msg.enemy_crit || Critical hit from {enemy}! {dmg} damage! || {enemy}의 치명타! {dmg} 피해!
combat.msg.enemy_hit || = || {enemy}이(가) {dmg}의 피해를 입혔습니다!
combat.msg.enemy_appeared || = || {enemy}이(가) 나타났습니다!
combat.msg.exp_gain || = || (경험치 +{exp})
combat.msg.level_up || = || 레벨 업 x{levels}!
combat.msg.max_level || = || 최대 레벨에 도달했습니다.
combat.msg.not_enough_mp || = || MP가 부족합니다.
combat.msg.player_couldnt_escape || = || 도망치지 못했습니다!
combat.msg.player_crit || = || 치명타! {dmg} 피해!
combat.msg.player_hit || = || {dmg}의 피해를 입혔습니다!
combat.msg.player_miss || = || 빗나갔습니다!
combat.msg.player_ran || = || 도망쳤습니다!
combat.msg.player_stunned || = || 기절했습니다!
combat.msg.skill.use.skill_rampage || = || 광폭해졌습니다!
combat.msg.skill_crit || = || 치명타 스킬! {dmg} 피해!
combat.msg.skill_hit || = || 스킬이 {dmg}의 피해를 입혔습니다!
combat.msg.skill_miss || = || 스킬이 빗나갔습니다!
combat.msg.skill_used || = || 스킬을 사용했습니다.
combat.msg.stat_point_one || = || 능력치 포인트 +{n}
combat.msg.stat_points_many || = || 능력치 포인트 +{n}
combat.msg.status_tick || = || {status}(으)로 {dmg}의 피해를 입었습니다.
combat.msg.status_tick_enemy || = || {status}(으)로 {enemy}이(가) {dmg}의 피해를 입었습니다.
combat.msg.stole_item || = || {item}을(를) 훔쳤습니다.
combat.msg.foresee || = || 예견: {action}
combat.msg.guard || = || 몸을 웅크립니다.
combat.msg.pushed_back_you || = || 차례가 뒤로 밀렸습니다!
combat.msg.windup.skill_vine_trap || = || {enemy}이(가) 덩굴을 감아 올립니다!
combat.msg.windup.skill_ghost_cut || = || {enemy}이(가) 칼날을 치켜듭니다!
combat.msg.windup_broken || = || {enemy}의 {skill}이(가) 끊겼습니다!
combat.msg.break || = || 브레이크! {name}이(가) 받는 피해 +50%!
combat.msg.steadfast || = || {name}은(는) 버티기 상태입니다!
combat.msg.steadfast_you || = || 버티기 상태입니다!
combat.foresee.stunned || = || 기절해서 아무것도 하지 못한다
combat.foresee.windup || = || {skill}을(를) 준비 중
# ---------------- cutscene: intro (narration, plain form) ----------------
cutscene.intro.0.0 || = || 이곳은 환상과 마법, 용의 세계가 아니다.
cutscene.intro.0.1 || = || 괴물은 없었다.
cutscene.intro.0.2 || = || 있는 것은 인간 속의 괴물뿐.
cutscene.intro.0.3 || This is a world of gritty medieval kingdoms and warring countries. || 이곳은 거칠고 투박한 중세 왕국과 전쟁 중인 나라들의 세계다.
cutscene.intro.0.5 || = || 정말 그럴까?
cutscene.intro.0.7 || You were a mapmaker in a remote village, scraping by penny to penny. || 당신은 외딴 마을의 지도 제작자였다. 한 푼 두 푼 겨우 모아 살았다.
cutscene.intro.0.8 || = || 두 왕국의 전쟁을 위해 징집이 있을 거라는 소문이 돌았다.
cutscene.intro.0.9 || The reward was enough to set you up for a year... || 보수는 1년을 먹고살기에 충분했다...
cutscene.intro.0.10 || = || 그만한 가치는 있었다...
cutscene.intro.0.11 || = || 당신은 지원했다.
cutscene.intro.1.0 || = || 계획은 지름길로 가서 적을 기습하는 것이었다.
cutscene.intro.1.1 || Good thing you grew up here and knew the forest. || 이곳에서 자라 숲을 잘 아는 게 다행이었다.
cutscene.intro.1.2 || = || 그러던 중 정찰병이 산 근처에서 터널을 발견했다고 보고했다.
cutscene.intro.1.3 || = || 평생 그런 터널은 들어 본 적이 없었다.
cutscene.intro.1.4 || = || 농담이라고 생각할 수밖에 없었다...
cutscene.intro.1.5 || = || 하지만 농담이 아니었다.
cutscene.intro.1.6 || = || 터널은 정말로 있었다. 믿기 어려웠다.
cutscene.intro.1.7 || = || 일행은 적의 성을 향해 터널 깊숙이 들어갔다...
cutscene.intro.1.8 || You were just the navigator for this raid party. || 당신은 이 토벌대의 길잡이일 뿐이었다.
cutscene.intro.1.9 || = || 그리고 이곳은 전에는 없던 곳이었다...
cutscene.intro.1.10 || = || 잘못된 일이라는 걸 알았다. 느낄 수 있었다.
cutscene.intro.1.11 || = || 입을 열기도 전에.
cutscene.intro.1.12 || = || 갈라지는 소리가 들린다...
cutscene.intro.1.13 || = || 작은 떨림...
cutscene.intro.2.0 || = || 하지만 이미 늦었다.
cutscene.intro.2.1 || = || 당신은 떨어졌다.
# ---------------- cutscene: ending (text kept, D104; Korean only) ----------------
cutscene.ending.0.0 || = || 신이시여, 감사합니다...
cutscene.ending.0.1 || = || 출구다...
cutscene.ending.0.2 || = || 여기가 어딘지 알아내야 해.
cutscene.ending.1.0 || = || 응?
cutscene.ending.1.2 || = || 저건... 토벌대 아닌가?
cutscene.ending.1.3 || = || 모두... 죽었어.
cutscene.ending.1.4 || = || ...대체 이게 무슨...
cutscene.ending.2.1 || = || 아.
cutscene.ending.2.3 || = || 결국 나는 운이 좋았던 게 아니었어.
cutscene.ending.3.0 || = || 희망은 없다.
cutscene.ending.4.0 || = || 플레이해 주셔서 감사합니다...
cutscene.ending.4.3 || = || 다음 작품도 응원해 주세요 :)
cutscene.game_over.0.0 || = || 게임 오버.
# ---------------- dialogue (inner speech: casual form; world text: plain form) ----------------
dialogue.after_class_chest_archer.0 || = || 가볍다. 균형이 좋다.
dialogue.after_class_chest_archer.1 || = || 녹도, 닳은 흔적도 없다.
dialogue.after_class_chest_archer.2 || = || 마치 여기서 기다리고 있었던 것처럼.
dialogue.after_class_chest_knight.0 || = || 무겁다... 그래도 몸에 맞아.
dialogue.after_class_chest_knight.1 || = || 몇 년 동안 내 것이었던 것처럼.
dialogue.after_class_chest_knight.2 || = || 이 상자가 나를 알 리 없는데.
dialogue.after_class_chest_mage.0 || = || 왜 이렇게 차갑지...?
dialogue.after_class_chest_mage.1 || = || ...방금 뭔가 속삭이지 않았나?
dialogue.after_class_chest_mage.2 || = || 아니야. 그럴 리 없어.
dialogue.after_class_chest_mage.3 || = || 나를 부르지 마.
dialogue.after_skillbook_chest_archer.0 || = || 각도와 호흡 표시.
dialogue.after_skillbook_chest_archer.1 || = || 첫 동작을 따라 해 보면...
dialogue.after_skillbook_chest_archer.2 || = || 나머지는 몸이 알아서 한다. 이상하다.
dialogue.after_skillbook_chest_knight.0 || = || 도해... 단순한 자세들.
dialogue.after_skillbook_chest_knight.1 || = || 한 번 읽었을 뿐인데.
dialogue.after_skillbook_chest_knight.2 || = || 어쩐지... 손이 이미 알고 있어.
dialogue.after_skillbook_chest_mage.0 || = || 읽을 수가 없어...
dialogue.after_skillbook_chest_mage.1 || = || ...페이지가 온통 기호투성이야.
dialogue.after_skillbook_chest_mage.2 || = || 어떤 목소리가 내 생각 속으로 스며든다.
dialogue.after_skillbook_chest_mage.3 || = || 그러자 이해가 된다. 손을 어디에 둬야 하는지 알려 준다.
dialogue.barrel_break.0 || = || 통을 부수었다.
dialogue.barrel_broken.0 || = || 이미 부서졌다.
dialogue.bone.0 || = || 오래된 뼈 무더기.
dialogue.chest_empty.0 || = || 텅 비어 있다.
dialogue.chest_open.0 || = || 상자를 열었다.
dialogue.dung_pile.0 || = || 손대지 않는 게 좋겠다.
dialogue.entering_second_floor.0 || = || 이해할 수 없어.
dialogue.entering_second_floor.1 || = || 이런 생물이 어떻게 존재하지?
dialogue.entering_second_floor.2 || = || 돌연변이 짐승 같은 건가...
dialogue.fire_stand_extinguish.0 || = || 불을 껐다.
dialogue.fire_stand_off.0 || = || 불이 꺼져 있다.
dialogue.first_encounter_slime.0 || = || 뭐...
dialogue.first_encounter_slime.1 || = || ...저게 대체 뭐야?
dialogue.horned_skull.0 || = || 그것이 당신을 마주 본다.
dialogue.loot_empty.0 || = || 아무것도 찾지 못했다.
dialogue.loot_received.0 || = || {item} x{qty}을(를) 받았습니다.
dialogue.rock.0 || = || 그냥 바위다.
dialogue.sys_class_chest_prompt.0 || = || 마침 잘됐군...
dialogue.sys_class_chest_prompt.1 || = || 여기서 나가려면 꼭 필요한 것이다.
dialogue.sys_floor1_intro.0 || = || *쿵* *와장창*
dialogue.sys_floor1_intro.1 || = || 뭐... 방금 무슨 일이 일어난 거야?
dialogue.sys_floor1_intro.2 || = || 여긴 어디지?
dialogue.sys_floor1_intro.4 || I knew it. There's something wrong about this place. || 역시 그랬어. 이곳은 뭔가 이상해.
dialogue.sys_floor1_intro.5 || = || 돌아가야 해.
dialogue.sys_load_ok.0 || = || 게임을 불러왔습니다.
dialogue.sys_save_ok.0 || = || 게임을 저장했습니다.
dialogue.torch_extinguish.0 || = || 횃불을 껐다.
dialogue.torch_off.0 || = || 횃불이 식어 있다.
dialogue.tree_stump.0 || = || 베어 낸 나무 그루터기.
# ---------------- enemies ----------------
enemy.name.enemy_direwolf || = || 다이어울프
enemy.name.enemy_final_boss || = || 창백한 성가대 짐승
enemy.name.enemy_ghostsword || = || 유령검
enemy.name.enemy_killer_plant || = || 살인 식물
enemy.name.enemy_madwhisp || = || 미친 도깨비불
enemy.name.enemy_mini_boss || = || 창백뿔
enemy.name.enemy_slime || = || 슬라임
enemy.name.enemy_stranger || = || 낯선 자
enemy.name.unknown || = || 적
# ---------------- world objects ----------------
interact.grave.default || = || 오래된 무덤이 있다.
interact.name.obj_barrel || = || 통
interact.name.obj_barrel_broken || = || 깨진 통
interact.name.obj_chest || = || 상자
interact.name.obj_chest_opened || = || 상자
interact.name.obj_fire_stand || = || 화로
interact.name.obj_fire_stand_off || = || 화로
interact.name.obj_horned_skull || = || 뿔 달린 해골
interact.name.obj_skull_1 || = || 해골
interact.name.obj_skull_2 || = || 해골
interact.name.obj_tall_grass || = || 긴 풀
interact.name.obj_torch || = || 횃불
interact.name.obj_torch_off || = || 횃불
interact.name.obj_tree_stump1 || = || 그루터기
interact.name.obj_tree_stump2 || = || 그루터기
interact.grave.line.obj_grave.0 || A worn grave marker sinks into the soil. || 닳은 묘비가 흙 속으로 가라앉고 있다.
interact.grave.line.obj_grave.1 || An old grave rests here, quiet and cold. || 오래된 무덤이 조용하고 차갑게 누워 있다.
interact.grave.line.obj_grave.2 || The stone is cracked, but the name is long gone. || 돌에는 금이 갔고, 이름은 오래전에 지워졌다.
interact.grave.line.obj_grave.3 || You brush away dust from a weathered grave. || 비바람에 깎인 무덤의 먼지를 털어 낸다.
interact.grave.line.obj_grave.4 || A faded grave stands in the silence. || 빛바랜 무덤이 침묵 속에 서 있다.
interact.grave.line.obj_grave.5 || Moss creeps over the edge of this grave. || 무덤 가장자리를 이끼가 뒤덮고 있다.
interact.grave.line.obj_grave.6 || The earth around this grave feels recently disturbed. || 이 무덤 주위의 흙은 최근에 파헤쳐진 듯하다.
interact.grave.line.obj_grave.7 || A neglected grave marker leans to one side. || 돌보는 이 없는 묘비가 한쪽으로 기울어 있다.
interact.grave.line.obj_grave.8 || This grave has been here for a long, long time. || 이 무덤은 아주 오래전부터 여기 있었다.
interact.grave.line.obj_grave.9 || Only silence answers from this old grave. || 이 오래된 무덤에서는 침묵만이 돌아온다.
# ---------------- items ----------------
item.msg.already_know || You already know {skill}. || 이미 {skill}을(를) 알고 있습니다.
item.msg.cant_use || = || 사용할 수 없습니다.
item.msg.class_cannot_learn || = || 현재 직업은 {skill}을(를) 배울 수 없습니다.
item.msg.full_health || HP is already full. || HP가 이미 가득 찼습니다.
item.msg.full_mana || MP is already full. || MP가 이미 가득 찼습니다.
item.msg.healed_hp || Restored {amt} HP. || HP를 {amt} 회복했습니다.
item.msg.learned || = || {skill}을(를) 배웠습니다.
item.msg.no_status_cure || = || 치료할 상태 이상이 없습니다.
item.msg.not_ready_skill || = || 아직 이 스킬을 배울 수 없습니다.
item.msg.recovered_mp || Restored {amt} MP. || MP를 {amt} 회복했습니다.
item.msg.status_cured || = || {status} 상태가 치료되었습니다.
item.msg.that_skill || = || 그 스킬
item.msg.throw_in_battle || = || 전투 중 적에게 던지세요.
item.name.0 || = || 없음
item.name.4 || = || 나무 활
item.name.5 || = || 철 활
item.name.6 || = || 백금 활
item.name.10 || = || HP 물약 Lv1
item.name.11 || = || MP 물약 Lv1
item.name.14 || = || HP 물약 Lv2
item.name.15 || = || HP 물약 Lv3
item.name.17 || = || MP 물약 Lv2
item.name.18 || = || MP 물약 Lv3
item.name.21 || = || 철 갑옷
item.passive.desc.1 || A Weak press counts as OK (once per turn) || 약함 판정이 무난으로 처리됩니다 (턴당 1회)
item.passive.desc.2 || First hit: -1 DEF for 2 turns || 첫 타격: 2턴 동안 DEF -1
item.passive.desc.3 || +35% damage to foes below 35% HP || HP 35% 미만 적에게 피해 +35%
item.passive.desc.4 || A Good press counts as Perfect (once per turn) || 좋음 판정이 완벽으로 처리됩니다 (턴당 1회)
item.passive.desc.5 || Good or Perfect: may Stun (3-turn cooldown) || 좋음 이상: 기절시킬 수 있음 (재사용 3턴)
item.passive.desc.6 || +30% damage to foes with a status || 상태 이상에 걸린 적에게 피해 +30%
item.passive.desc.7 || = || 2턴마다 MP +1
item.passive.desc.8 || = || 스킬 피해 +15%, 상태 지속 +1턴
item.passive.desc.9 || +25% skill damage, Perfect +1 MP || 스킬 피해 +25%, 완벽 판정 시 MP +1
item.passive.desc.20 || First hit -3 damage (once per battle) || 첫 피격 피해 -3 (전투마다)
item.passive.desc.21 || After a hit: next hit -40% (2-turn cooldown) || 피격 후 다음 피격 피해 -40% (재사용 2턴)
item.passive.desc.22 || -30% damage from big hits || 큰 피해 30% 감소
item.passive.desc.23 || +1 MP every 2 turns, +1 physical damage taken || 2턴마다 MP +1, 받는 물리 피해 +1
item.passive.desc.24 || = || 스킬 피해 +10%, 상태 지속 +1턴
item.passive.desc.25 || Reflects the first status (costs 1 MP) || 첫 상태 이상을 반사 (MP 1 소모)
item.passive.desc.44 || = || 받는 피해 10% 감소
item.passive.desc.45 || = || 3턴마다 MP +1
item.passive.desc.46 || = || 골드를 15% 더 얻음
item.passive.name.1 || = || 안정된 일격
item.passive.name.2 || = || 분쇄의 칼날
item.passive.name.3 || = || 처형자
item.passive.name.4 || = || 조준 보조
item.passive.name.5 || = || 속박 사격
item.passive.name.6 || = || 포식자
item.passive.name.7 || = || 마나 방울
item.passive.name.8 || = || 불씨 렌즈
item.passive.name.9 || = || 비전 쇄도
item.passive.name.20 || = || 첫 충격
item.passive.name.21 || = || 대비
item.passive.name.22 || = || 빚의 판금
item.passive.name.23 || = || 얇은 장막
item.passive.name.24 || = || 비단의 흐름
item.passive.name.25 || = || 거울 바느질
item.passive.name.44 || = || 단단함
item.passive.name.45 || = || 샘물
item.passive.name.46 || = || 행운
item.sent.30 || = || 위력 {power}의 피해를 주고 {turns}턴 동안 불탑니다. 턴을 소모합니다.
item.sent.accessory || = || 장신구 칸에 착용하는 반지나 부적입니다.
item.sent.accessory_stat || = || 착용 효과: {bonus}.
item.sent.accessory_passive || = || 착용 효과: {line}.
item.sent.mp || = || MP를 {min}에서 {max}만큼 회복합니다.
item.sent.cure || = || {status} 상태를 치료합니다.
item.sent.armor || = || 방어구: {bonus}.
# ---------------- menus ----------------
menu.common.confirm || = || 확인
menu.common.yes || = || 예
menu.common.no || = || 아니요
menu.difficulty.title || Select difficulty || 난이도 선택
menu.label.equipped_short || = || 장
menu.main.option.0 || New game || 새 게임
menu.main.option.1 || Load game || 게임 불러오기
menu.main.option.2 || = || 설정
menu.main.option.3 || Exit game || 게임 종료
menu.main.option.continue || = || 이어하기
menu.msg.already_equipped || = || 이미 장착 중입니다.
menu.msg.cant_equip || = || 장착할 수 없습니다.
menu.msg.equipped || = || {item}을(를) 장착했습니다.
menu.msg.no_items || = || 아이템이 없습니다.
menu.msg.no_skills || = || 스킬이 없습니다.
menu.msg.not_in_inventory || = || 소지품에 없는 아이템입니다.
menu.msg.unequipped || = || {item}을(를) 해제했습니다.
menu.popup.learn_skill || = || 스킬 배우기
menu.stats.available_points || Available points: {points} || 사용 가능한 포인트: {points}
menu.stats.level_cap || Level cap || 최대 레벨
menu.stats.required_exp || = || 필요 경험치
menu.tab.inventory || = || 소지품
menu.tab.stats || = || 능력치
menu.tooltip.class.all || All classes || 모든 직업
menu.tooltip.class.only || = || {classes} 전용
menu.tooltip.damage || = || 피해
menu.tooltip.damage.acc || = || 명중 {acc}
menu.tooltip.damage.hits || = || {hits}회
menu.tooltip.damage.power || = || 위력 {power}
menu.tooltip.damage.stat || = || 능력치 {stat}
menu.tooltip.effect.buff || = || 효과: 버프
menu.tooltip.effect.multi_status || = || 효과: 여러 상태 이상
menu.tooltip.effect.restore_hp || = || 효과: HP 회복
menu.tooltip.effect.steal_item || = || 효과: 무작위 아이템 1개를 훔칩니다.
menu.tooltip.item.cure || = || 치료: {status}
menu.tooltip.item.heal_hp || Restore HP {min}-{max} || HP 회복 {min}-{max}
menu.tooltip.item.restore_mp || = || MP 회복 {min}-{max}
menu.tooltip.item.scale_with_level || = || 수치는 레벨에 따라 달라집니다.
menu.tooltip.item.teaches || = || 습득 스킬: {skill}
menu.tooltip.item.teaches_unknown || = || 습득 스킬: 알 수 없음
menu.tooltip.item.type.armor || = || 방어구
menu.tooltip.item.type.key || Key item || 중요 아이템
menu.tooltip.item.type.unknown || = || 알 수 없음
menu.tooltip.item.type.with_slot || = || {type} (부위: {slot})
menu.tooltip.passive.cooldown || = || 재사용 대기
menu.tooltip.passive.cooldown_prefix || = || (재사용 대기
menu.tooltip.passive.good_plus || Good or Perfect || 좋음 또는 완벽
menu.tooltip.passive.good_plus_colon || Good or Perfect: || 좋음 또는 완벽:
menu.tooltip.passive.once_per_turn || = || 턴당 1회
menu.tooltip.skill.mp_target || = || MP {mp} | 대상: {target}
menu.tooltip.skill.target_suffix || = || | 대상: {target}
menu.tooltip.stat.spd.desc || = || 누가 먼저 행동할지 정합니다.
menu.tooltip.stat.def.desc || = || 받는 모든 피해를 줄입니다.
menu.tooltip.stat.default.desc || = || 전투 능력에 영향을 줍니다.
menu.tooltip.stat.mag.desc || = || 마법 피해와 최대 MP를 늘립니다.
menu.tooltip.stat.atk.desc || = || 물리 피해(기본 공격과 ATK 스킬)를 늘립니다.
menu.tooltip.status.turns_short || = || {turns}턴
menu.tooltip.status.turns_suffix || = || | {turns}턴
menu.tooltip.statuses.more || = || , 외 {count}개
menu.tooltip.target.all_allies || All allies || 아군 전체
menu.tooltip.target.all_enemies || All enemies || 적 전체
menu.tooltip.target.enemy || Enemy || 적
menu.tooltip.target.self || = || 자신
menu.tooltip.target.unknown || = || 알 수 없음
menu.tooltip.item.type.accessory || = || 장신구
pause.option.exit_main || Exit to main menu || 메인 메뉴로 나가기
pause.option.resume || = || 계속하기
pause.option.abandon || Abandon run || 도전 포기
pause.option.abandon_confirm || = || 정말입니까? 한 번 더 확인
# ---------------- save ----------------
save.back || = || 뒤로
save.confirm.delete || = || 이 슬롯을 삭제하시겠습니까?
save.confirm.load || = || 이 저장 데이터를 불러오시겠습니까?
save.confirm.overwrite || = || 덮어쓰시겠습니까?
save.msg.no_game_saves || No saved games found. || 저장된 게임이 없습니다.
save.msg.no_saved_games_found || = || 저장된 게임이 없습니다.
save.level.short || = || Lv
save.title.load || Load game || 게임 불러오기
save.title.save || Save game || 게임 저장
# ---------------- settings ----------------
settings.row.bgm || Music || 음악
settings.row.sfx || Effects || 효과음
settings.row.ui || Interface || 인터페이스
settings.row.scale || = || 화면 배율
settings.row.qte_assist || Timing assist || 타이밍 보조
settings.section.audio || = || 소리
settings.section.display || = || 화면
settings.value.pending || = || 대기 중
settings.value.speed_slow || = || 느림
settings.value.speed_fast || = || 빠름
settings.help.qte_assist || = || 막대 판정 구간이 넓어지고 화살표 스킬의 시간이 늘어납니다.
# ---------------- skills ----------------
skill.name.skill_bite || = || 깨물기
skill.name.skill_bloodthirsty || = || 피에 굶주림
skill.name.skill_double_shot || = || 연속 사격
skill.name.skill_fireball || = || 화염구
skill.name.skill_ghost_cut || = || 베기
skill.name.skill_hilt_bash || = || 자루 강타
skill.name.skill_horiz_slash || = || 가로 베기
skill.name.skill_ice_spear || = || 얼음 창
skill.name.skill_meditation || = || 명상
skill.name.skill_muscle_up || = || 근육 강화
skill.name.skill_poison_fangs || = || 독니
skill.name.skill_poison_mist || = || 독안개
skill.name.skill_ramming || = || 돌진
skill.name.skill_rampage || = || 날뛰기
skill.name.skill_rev_up || = || 기세 상승
skill.name.skill_scorching_touch || = || 작열의 손길
skill.name.skill_solidify || = || 굳히기
skill.name.skill_steal || = || 훔치기
skill.name.skill_vine_trap || = || 덩굴 덫
skill.name.skill_wound || = || 상처 내기
skill.name.skill_regrow || = || 다시 자라기
skill.name.skill_snuff || = || 꺼뜨리기
skill.name.skill_gore || = || 들이받기
skill.name.unknown || = || 알 수 없음
item.name.101 || = || 스킬북: 상처 내기
item.name.102 || = || 스킬북: 자루 강타
item.name.103 || = || 스킬북: 근육 강화
item.name.104 || = || 스킬북: 기세 상승
item.name.105 || = || 스킬북: 가로 베기
item.name.106 || = || 스킬북: 독화살
item.name.107 || = || 스킬북: 회피
item.name.108 || = || 스킬북: 연속 사격
item.name.109 || = || 스킬북: 조준
item.name.110 || = || 스킬북: 화염구
item.name.111 || = || 스킬북: 독안개
item.name.112 || = || 스킬북: 얼음 창
item.name.113 || = || 스킬북: 명상
item.name.114 || = || 스킬북: 선견
# ---------------- statuses ----------------
status.name.stat.spd || = || SPD
status.name.stat.def || = || DEF
status.name.stat.mag || = || MAG
status.name.stat.atk || = || ATK
status.name.status_bleed || Bleed || 출혈
status.name.status_burn || Burn || 화상
status.name.status_bloodthirsty || = || 피에 굶주림
status.name.status_dmg_up || = || 기세 상승
status.name.status_guard || = || 방어
status.name.status_meditation || = || 명상
status.name.status_solidify || = || 굳히기
status.name.status_stun || = || 기절
status.name.status_sunder || = || 분쇄
status.name.status_brace || = || 방어 태세
status.name.unknown || = || 알 수 없음
status.desc.status_stun || = || 차례가 뒤로 밀립니다. 이후 2턴간 버티기 상태가 되어 기절하지 않습니다.
status.desc.status_hit_up || = || {turns}턴 동안 적의 다음 행동이 보이고 모든 완벽 구간이 {pct}% 넓어집니다.
# ---------------- elites, throwables, shop ----------------
elite.line.armored || = || 철갑: DEF +{n}.
elite.line.swift || = || 신속: SPD +{n}.
elite.line.brutal || = || 흉포: ATK +{n}.
throw.msg.burn || It burns for {n} each turn. || 매 턴 {n}의 화상을 입습니다.
throw.msg.thrown || = || {item}을(를) 던졌습니다!
shop.lbl.stat || = || 기준 능력치
shop.val.pushback || = || 밀치기
shop.val.escape || = || 도망
shop.lbl.nobody || = || 무명
shop.lbl.perfect || = || 완벽
shop.msg.equipped || = || 먼저 장착을 해제하세요.
shop.msg.slots || = || 스킬 칸이 가득 찼습니다.
# ---------------- cards ----------------
card.lbl.perfect || = || 완벽 구간
card.pas.venom_coat.line || Poison you apply adds {n} extra stack. || 내가 거는 독이 {n}중첩 더 쌓인다.
card.pas.field_medic.line || Restore {n} HP after each fight you win. || 전투에서 이길 때마다 HP를 {n} 회복한다.
card.pas.k_iron_hide.line || Thicker hide: +{n} DEF for good. || 두꺼운 가죽: 영구히 DEF +{n}.
card.pas.potion_lore.name || = || 물약 지식
class.nobody.locked_hint || = || 한 번 승리하면 이 길이 열립니다.
# ---------------- tutorial ----------------
tutorial.stagger.title || Stagger and BREAK || 경직과 브레이크
tutorial.stagger.body || Good hits fill the enemy's stagger meter. When it is full the enemy goes into BREAK: it loses a turn and takes +50% damage. || 좋은 공격은 적의 경직 게이지를 채웁니다. 가득 차면 적이 브레이크 상태가 되어 한 턴을 잃고 피해를 50% 더 받습니다.
tutorial.status.body || Poison, Burn, Bleed and Stun last a few turns. The icons under a name show which ones a fighter has. || 독, 화상, 출혈, 기절은 몇 턴 동안 이어집니다. 이름 아래 아이콘이 걸려 있는 상태를 보여 줍니다.
# ---------------- death scene, wake, notes (bleak narration, plain form) ----------------
death.killer.enemy_slime || = || 슬라임. 고작 슬라임에게 당했다.
note.wall_breath.0 || = || 돌에 깊이, 천천히, 조심스럽게 새겨져 있다.
note.carved_arrow.2 || = || "가지 마."
# ---------------- credits ----------------
credits.creator || = || 제작: {name}
credits.font_galmuri || = || Galmuri - Lee Minseo (OFL-1.1)
"""

# New rows (ref || en || ko). Updated in place when the ref already exists.
ADDS = """
combat.msg.passive.mirror_stitch || Mirror Stitch repelled a status! || 거울 바느질이 상태 이상을 튕겨 냈습니다!
combat.msg.passive.pinning_shot || Pinning Shot stunned the enemy! || 속박 사격이 적을 기절시켰습니다!
combat.msg.defence.resisted || Resisted {status}. || {status}을(를) 견뎌 냈습니다.
combat.msg.defence.summary || {grade} block! Damage -{reduce}%. || {grade} 방어! 피해 -{reduce}%.
combat.msg.defence.weak || Not blocked. || 막지 못했습니다.
combat.msg.player_cant_escape_this || You can't run away from this. || 이 상대에게서는 도망칠 수 없습니다.
merchant.greet.0 || You came a long way down. Brave. Or lost. || 꽤 깊이 내려왔군. 용감한 건가, 길을 잃은 건가.
merchant.greet.1 || Coins keep the dark away. Mostly. || 동전은 어둠을 쫓아 주지. 대개는.
merchant.greet.2 || Don't ask where I sleep. || 내가 어디서 자는지는 묻지 마.
merchant.greet.3 || Buy something. Please. || 뭐라도 사. 제발.
merchant.buy.0 || Good choice. You will need it. || 좋은 선택이야. 필요하게 될 거다.
merchant.buy.1 || It belonged to someone once. || 한때는 누군가의 물건이었지.
merchant.buy.2 || Spend it while you can. || 쓸 수 있을 때 써 둬.
merchant.sell.0 || I'll take it. No questions. || 받아 두지. 묻진 않을게.
merchant.sell.1 || It will find its way to someone else. || 그건 또 누군가에게 흘러가겠지.
merchant.sell.2 || Used. Like everything down here. || 낡았군. 여기 있는 건 다 그렇지.
merchant.leave.0 || Come back. They always do. || 또 와. 다들 그러더군.
merchant.leave.1 || Mind the dark. It minds you. || 어둠을 조심해. 어둠도 널 지켜보니까.
merchant.leave.2 || I'll be here. I'm always here. || 난 여기 있을 거야. 늘 여기 있어.
"""

# Global term normalisation, applied to every row after the edits (one name per thing).
KO_TERMS = [
    ("기술", "스킬"),
    ("포션", "물약"),
    ("스탯", "능력치"),
    ("퍼펙트", "완벽"),
    ("마술사", "마법사"),
]
EN_TERMS = [
    (re.compile(r"\bdefence\b"), "defense"),
    (re.compile(r"\bDefence\b"), "Defense"),
    (re.compile(r"\barmour\b"), "armor"),
    (re.compile(r"\bArmour\b"), "Armor"),
]


def parse(block: str):
    for raw in block.splitlines():
        line = raw.rstrip("\n")
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        parts = line.split(" || ")
        if len(parts) != 3:
            raise SystemExit(f"bad line (need 'ref || en || ko'): {line!r}")
        yield parts[0].strip(), parts[1], parts[2]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=str(XLSX))
    args = ap.parse_args()
    path = Path(args.xlsx)

    wb = load_workbook(path)
    ws = wb["strings"]
    header = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
    cref, cen, cko = header.index("ref") + 1, header.index("en") + 1, header.index("ko") + 1

    index = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, cref).value
        if v is not None:
            index[str(v)] = r

    # 1. junk line-number rows + surplus grave rows
    kill = []
    for ref, r in index.items():
        m = JUNK_RE.match(ref)
        if m and m.group("stem") in JUNK_STEMS:
            kill.append(r)
        elif ref.startswith("interact.grave.line.obj_grave."):
            n = int(ref.rsplit(".", 1)[1])
            if n >= 10:
                kill.append(r)
    for r in sorted(set(kill), reverse=True):
        ws.delete_rows(r)
    index = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, cref).value
        if v is not None:
            index[str(v)] = r

    changed_en = changed_ko = added = 0
    missing = []

    def put(ref, en, ko, create):
        nonlocal changed_en, changed_ko, added
        r = index.get(ref)
        if r is None:
            if not create:
                missing.append(ref)
                return
            r = ws.max_row + 1
            ws.cell(r, cref).value = ref
            index[ref] = r
            added += 1
        if en != "=" and (ws.cell(r, cen).value or "") != en:
            ws.cell(r, cen).value = en
            changed_en += 1
        if ko != "=" and (ws.cell(r, cko).value or "") != ko:
            ws.cell(r, cko).value = ko
            changed_ko += 1

    for ref, en, ko in parse(EDITS):
        put(ref, en, ko, create=False)
    for ref, en, ko in parse(ADDS):
        put(ref, en, ko, create=True)

    # the grave rows exist for 0-9 only if the sheet had them; make sure all ten are there
    for ref, en, ko in parse(EDITS):
        if ref.startswith("interact.grave.line.obj_grave.") and ref in missing:
            put(ref, en, ko, create=True)
            missing.remove(ref)

    # 3. global terms
    for r in range(2, ws.max_row + 1):
        en = ws.cell(r, cen).value
        ko = ws.cell(r, cko).value
        if isinstance(en, str):
            new = en
            for rx, rep in EN_TERMS:
                new = rx.sub(rep, new)
            if new != en:
                ws.cell(r, cen).value = new
                changed_en += 1
        if isinstance(ko, str):
            new = ko
            for a, b in KO_TERMS:
                new = new.replace(a, b)
            if new != ko:
                ws.cell(r, cko).value = new
                changed_ko += 1

    wb.save(path)
    print(f"deleted {len(set(kill))} rows, added {added}, en cells changed {changed_en}, ko cells changed {changed_ko}")
    if missing:
        print("refs not in the workbook (skipped):", ", ".join(missing))
    return 0


if __name__ == "__main__":
    sys.exit(main())
