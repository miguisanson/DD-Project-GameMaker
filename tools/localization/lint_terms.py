#!/usr/bin/env python3
"""Check the wording of the UI text (D133): one word per meaning, Title Case for labels. Exit 1 when something is wrong.

The chosen words and the variants they replaced are in docs/terminology-g11.md; this file holds the same table as data.
Flagged, per workbook row (sheet `strings`):
  - a banned English variant ("Foes", "Take off", "Exit", "Select", "XP", "health" ...) in a row that is in scope;
  - a banned Korean variant (착용, 상대, 떠나기, 차례 ...) in a row that is in scope;
  - an English label that is not Title Case. Which rows are labels is decided by key (LABEL_KEYS below): tabs, menu rows, buttons,
    headers, titles, row/stat labels, names and the hint words next to a button glyph. Sentences (messages, help lines, descriptions) stay sentence case.

"In scope" = every row except the voice of the story: dialogue, cutscenes, choices, notes, death and wake lines, credits, graves,
the merchant's own lines and the hero's first-fight reaction. Those keep their own words (a cutscene may say "monster").

Usage:
  python tools/localization/lint_terms.py [--xlsx datafiles/localization/game_text.xlsx]
  python tools/localization/lint_terms.py --selftest
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import List, Optional, Tuple

SHEET = "strings"

# Rows with the story's own voice: not checked for the banned variants.
STORY_PREFIXES = (
    "dialogue.", "cutscene.", "choice.", "note.", "death.", "wake.", "credits.",
    "interact.grave.", "merchant.greet.", "merchant.buy.", "merchant.sell.", "merchant.leave.", "combat.msg.floor1_reaction.",
)

# (concept, preferred, regex on the English text, keys exempt from this rule)
EN_BANNED: List[Tuple[str, str, str, Tuple[str, ...]]] = [
    ("enemy", "enemy / enemies", r"\b(foes?|monsters?|creatures?|opponents?)\b", ()),
    ("battle", "battle", r"\b(fights?|combat)\b", ()),
    ("HP / MP", "HP / MP", r"\b(health|hit points?|magic points?)\b", ("help.line.hp", "help.line.mp")),
    ("restore", "restore", r"\b(heals?|healed|recovers?|recovered)\b", ()),
    ("quit", "Quit", r"\bexit(s|ed|ing)?\b", ()),
    ("title", "title", r"\bmain menu\b", ()),
    ("unequip", "Unequip", r"\btake (it )?off\b", ()),
    ("equipped", "Equipped", r"\b(worn|wear|wears|wearing)\b", ()),
    ("choose", "Choose", r"\bselect(s|ed)?\b", ()),
    ("abandon", "Abandon", r"\b(gave up|gave|give up|gives up|giving up)\b", ()),
    ("continue", "Continue", r"\bgo on\b", ()),
    ("back", "Back", r"\breturn\b", ()),
    ("EXP", "EXP", r"\b(xp|experience)\b", ()),
    ("evasion", "Evasion", r"^\s*dodge\s*$", ()),
    ("bag", "Bag", r"\b(inventory|backpack)\b", ()),
]

# (concept, preferred Korean, regex on the Korean text, keys exempt)
KO_BANNED: List[Tuple[str, str, str, Tuple[str, ...]]] = [
    ("equip", "장착", r"착용", ()),
    ("enemy", "적", r"상대(?!적)|몬스터|적군", ()),
    ("leave", "나가기", r"떠나기", ()),
    ("rest", "휴식", r"쉬기", ()),
    ("choose", "선택", r"고르기", ("card.fb.rest.name",)),   # 숨 고르기 is an idiom (catch your breath)
    ("take", "받기", r"가져가기", ()),
    ("delete", "삭제", r"지우기", ()),
    ("turn", "턴", r"차례", ()),
    ("bag", "소지품", r"가방|인벤토리", ()),
    ("stat", "능력치", r"스탯", ()),
    ("potion", "물약", r"포션", ()),
    ("damage", "피해", r"데미지", ()),
    ("crit", "치명타", r"크리티컬", ()),
    ("perfect", "완벽", r"퍼펙트", ()),
    ("class", "직업", r"수업", ()),
    ("run", "도전", r"모험", ()),
    ("battle", "전투", r"싸움", ()),
    ("HP / MP", "HP / MP", r"체력|마력", ("help.line.hp", "help.line.mp")),
    ("beaten", "처치한 적", r"쓰러뜨린 적", ()),
]

# Rows whose English is a label (Title Case). Matched against the whole key.
LABEL_KEYS = [
    r"battle\.action\.[a-z_.]+", r"battle\.menu\.back", r"battle\.det\.[a-z]+", r"battle\.hud\.[a-z]+", r"battle\.preview\.[a-z]+",
    r"battle\.chip\.steadfast", r"battle\.hint\.list",
    r"bed\.option\.[a-z_]+", r"bed\.title", r"bed\.lbl\.[a-z]+", r"bed\.upgrade_title",
    r"card\.lbl\.[a-z_]+", r"card\.tag\.[a-z_]+", r"card\.pas\.[a-z_]+\.name", r"card\.fb\.[a-z_]+\.name",
    r"cards\.pick\.(title|hint|skip)", r"cards\.stat\.(title|after|now|maxhp|maxmp)", r"cards\.replace\.(title|keep|lose)",
    r"class\.select\.(title|become)", r"class\.name\.[a-z_]+", r"class\.nobody\.locked", r"class\.lbl\.[a-z]+",
    r"common\.[a-z]+", r"controls\.(title|row\.[a-z_]+|reset)",
    r"echo\.pick\.title", r"elite\.(tag|name\.[a-z]+)", r"enemy\.name\.[a-z_]+",
    r"help\.title", r"help\.stats\.title", r"help\.term\.[a-z]+",
    r"interact\.name\.[a-z_0-9]+", r"item\.name\.[0-9]+", r"item\.passive\.name\.[0-9]+",
    r"menu\.common\.[a-z]+", r"menu\.difficulty\.title", r"menu\.main\.option\.([0-9]+|continue)", r"menu\.tab\.[a-z]+", r"menu\.popup\.[a-z_]+",
    r"menu\.gold", r"menu\.stats\.(level|level_cap|max|required_exp|available_points)",
    r"menu\.tooltip\.(class\.(all|only)|damage(\.(acc|power|stat))?|effect\.[a-z_]+|item\.(bonus|cure|heal_hp|restore_mp|teaches|teaches_unknown|type\.(armor|consumable|key|skillbook|unknown|weapon|accessory))|skill\.mp|target\.[a-z_]+|status\.(default|line)|statuses\.prefix)",
    r"merchant\.name", r"pause\.(option\.[a-z_]+|title)",
    r"runflow\.(sum\.[a-z]+|echo\.(item|skill|title|hint)|falls|floor|cause\.enemy)",
    r"save\.(back|delete|title\.[a-z]+|slot\.prefix|level\.short)",
    r"seed\.(choice\.[a-z]+|entry\.(title|hint[0-9]?))",
    r"settings\.(row\.[a-z_]+|section\.[a-z]+|title|value\.[a-z_]+|difficulty\.option\.[0-9]+|line\.[a-z_]+|keep\.(title|hint)|lang\.[a-z]+)",
    r"shop\.(lbl\.[a-z_]+|val\.(pushback|escape)|title|buy|sell|leave|sold_out|worn|gold|nothing_to_sell)",
    r"skill\.name\.[a-z_]+", r"stat\.[a-z_]+", r"status\.name\.[a-z_.]+",
    r"tab\.(lbl|val|title|hint)\.[a-z_]+", r"tab\.(codex|gear|items|skills|status|settings|system)",
    r"tab\.codex\.(bestiary|history|notes|unlocks|bestiary_title|notes_title)", r"tab\.gear\.(bag|worn|empty|no_[a-z]+)",
    r"tab\.skills\.(passives|none)", r"tab\.items\.none", r"tab\.unlock\.nobody", r"tab\.bestiary\.attack",
    r"title\.(card\.(cant|time)|splash\.[a-z_]+)", r"tutorial\.[a-z]+\.title", r"tutorial\.footer",
    r"ui\.skip_hold",
]
LABEL_RES = [re.compile(r"^" + p + r"$") for p in LABEL_KEYS]

# English labels that legitimately carry two Korean words (two meanings behind one English word).
SAME_EN_TWO_KO_OK = {"run", "ok", "no", "effects", "brace", "damage taken"}

SMALL = {"a", "an", "the", "of", "to", "in", "on", "at", "for", "and", "or", "but", "by", "with", "into", "from", "as", "vs", "per"}
CAPS_OK = {"BREAK", "PERFECT", "GOOD", "WEAK"}   # R3 note 9 (D171): the QTE grade words are the one all-capitals exception; OK is short enough to pass anyway


def in_scope(key: str) -> bool:
    return not key.startswith(STORY_PREFIXES)


def strip_markup(text: str) -> str:
    text = re.sub(r"\[k:[a-z_]+\]", " ", text)
    return re.sub(r"\[/?[ta]\]", "", text)


def title_case_problems(text: str) -> List[str]:
    """Words of an English label that break Title Case. Placeholders, digits, short abbreviations (HP, MP, OK) and BREAK pass."""
    problems: List[str] = []
    for seg in re.split(r"\s{2,}|\s\|\s", strip_markup(text)):
        words = seg.split()
        for i, w in enumerate(words):
            core = w.strip("()[]:;,.!?\"'+&/")
            if not core or not re.search(r"[A-Za-z]", core) or "{" in core or "}" in core or re.search(r"\d", core):
                continue
            core = re.split(r"[-/]", core)[0]
            if not core or not core[0].isalpha():
                continue
            first = i == 0 or words[i - 1].endswith(":")
            last = i == len(words) - 1
            if core.isupper():
                if len(core) > 3 and core not in CAPS_OK:
                    problems.append("'%s' is all caps" % core)
                continue
            if core.lower() in SMALL and not first and not last:
                if core[0].isupper():
                    problems.append("'%s' should be lower case" % core)
            elif core[0].islower():
                problems.append("'%s' should start with a capital" % core)
    return problems


def check_row(key: str, en: str, ko: str) -> List[str]:
    errs: List[str] = []
    if in_scope(key):
        for concept, preferred, rx, exempt in EN_BANNED:
            if key not in exempt:
                m = re.search(rx, en, re.I)
                if m:
                    errs.append("%s [en]: '%s' -> %s (%s)" % (key, m.group(0).strip(), preferred, concept))
        for concept, preferred, rx, exempt in KO_BANNED:
            if key not in exempt:
                m = re.search(rx, ko)
                if m:
                    errs.append("%s [ko]: '%s' -> %s (%s)" % (key, m.group(0), preferred, concept))
    if any(r.match(key) for r in LABEL_RES):
        for p in title_case_problems(en):
            errs.append("%s [en]: '%s' is not Title Case: %s" % (key, en, p))
    return errs


def consistency(rows: List[Tuple[str, str, str]]) -> List[str]:
    """One English label, one Korean word: the same label text on two label rows must have the same Korean."""
    seen: dict = {}
    errs: List[str] = []
    for key, en, ko in rows:
        if not en or not any(r.match(key) for r in LABEL_RES) or en.strip().lower() in SAME_EN_TWO_KO_OK:
            continue
        norm = en.strip().lower()
        if norm in seen and seen[norm][1] != ko:
            errs.append("%s [ko]: '%s' for '%s' but %s has '%s'" % (key, ko, en, seen[norm][0], seen[norm][1]))
        else:
            seen.setdefault(norm, (key, ko))
    return errs


def selftest() -> int:
    bad = 0

    def expect(name: str, got, want) -> None:
        nonlocal bad
        if got != want:
            bad += 1
            print("SELFTEST FAIL %s: got %r want %r" % (name, got, want))

    expect("title case ok", title_case_problems("Quit to Title"), [])
    expect("small word capitalised", title_case_problems("Quit To Title"), ["'To' should be lower case"])
    expect("lower word", title_case_problems("Sold out"), ["'out' should start with a capital"])
    expect("all caps word", title_case_problems("ATTACK"), ["'ATTACK' is all caps"])
    expect("abbreviations pass", title_case_problems("Max HP +{n}"), [])
    expect("hint words after a glyph", title_case_problems("[k:left][k:right] Choose   [k:confirm] take"), ["'take' should start with a capital"])
    expect("last word is capitalised", title_case_problems("Look Up"), [])
    expect("a colon starts a new clause", title_case_problems("Effect: buff"), ["'buff' should start with a capital"])
    expect("placeholders and digits pass", title_case_problems("{name} x{qty}") + title_case_problems("Level 3"), [])
    expect("banned en", check_row("tab.lbl.kills", "Foes Beaten", "처치한 적"), ["tab.lbl.kills [en]: 'Foes' -> enemy / enemies (enemy)"])
    expect("banned ko", check_row("tab.gear.worn", "Equipped", "착용 중"), ["tab.gear.worn [ko]: '착용' -> 장착 (equip)"])
    expect("story rows keep their words", check_row("cutscene.intro.3", "A monster, then a fight.", "괴물"), [])
    expect("sentences are not labels", check_row("menu.msg.equipped", "Equipped {item}.", "{item}을(를) 장착했습니다."), [])
    expect("label rows are checked", len(check_row("settings.row.back", "back", "뒤로")), 1)
    expect("exempt definition rows", check_row("help.line.hp", "Health. At 0 HP you fall.", "체력입니다."), [])
    expect("status chips are codes (D150), not labels", check_row("status.short.status_stun", "Stn", "기절"), [])
    expect("BREAK! is a term, not shouting", check_row("battle.hud.break", "BREAK!", "브레이크!"), [])
    expect("same label, same Korean", consistency([("bed.option.rest", "Rest", "휴식"), ("bed.option.heal", "Rest", "휴식")]), [])
    expect("same label, two Korean", consistency([("bed.option.rest", "Rest", "휴식"), ("bed.option.heal", "Rest", "쉬기")]),
           ["bed.option.heal [ko]: '쉬기' for 'Rest' but bed.option.rest has '휴식'"])
    expect("two meanings are allowed", consistency([("battle.action.run", "Run", "도망"), ("controls.row.run", "Run", "달리기")]), [])
    print("selftest %s" % ("FAILED" if bad else "ok"))
    return 1 if bad else 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default="datafiles/localization/game_text.xlsx")
    ap.add_argument("--selftest", action="store_true")
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")   # Korean in the messages, whatever the console code page
    if args.selftest:
        return selftest()
    from openpyxl import load_workbook

    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    ws = wb[SHEET]
    errs: List[str] = []
    rows = 0
    allrows: List[Tuple[str, str, str]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not row or not row[0]:
            continue
        key = str(row[0]).strip()
        en = "" if row[1] is None else str(row[1])
        ko = "" if row[2] is None else str(row[2])
        rows += 1
        allrows.append((key, en, ko))
        errs.extend(check_row(key, en, ko))
    errs.extend(consistency(allrows))
    for e in errs:
        print("TERMS " + e)
    print("terms-lint: %d row(s), %d problem(s)" % (rows, len(errs)))
    return 1 if errs else 0


if __name__ == "__main__":
    sys.exit(main())
