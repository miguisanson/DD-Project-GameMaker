#!/usr/bin/env python3
"""Verify localization workbook equivalence across schemas.

Compares effective runtime dictionaries (key -> {en, ko}) from two workbooks,
independent of legacy/minimal column layout.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Dict, Tuple

from openpyxl import load_workbook

SHEET_NAME = "strings"
LEGACY_REQUIRED = ("key", "english_en", "korean_ko")
MINIMAL_REQUIRED = ("ref", "en", "ko")


def norm(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def parse_active(value: object) -> int:
    s = norm(value)
    if s == "":
        return 1
    try:
        return 1 if int(float(s)) != 0 else 0
    except ValueError:
        return 0


def detect(headers: Dict[str, int]) -> str:
    if all(c in headers for c in LEGACY_REQUIRED):
        return "legacy"
    if all(c in headers for c in MINIMAL_REQUIRED):
        return "minimal"
    raise ValueError("Unrecognized workbook schema")


def get_val(values: Dict[str, object], schema: str, logical: str) -> str:
    if schema == "legacy":
        mapping = {
            "key": "key",
            "ref": "key",
            "en": "english_en",
            "ko": "korean_ko",
            "active": "active",
        }
    else:
        mapping = {
            "key": "key",
            "ref": "ref",
            "en": "en",
            "ko": "ko",
            "active": "active",
        }
    col = mapping[logical]
    if logical in {"en", "ko"}:
        out = text(values.get(col))
    else:
        out = norm(values.get(col))
    if schema == "minimal" and logical == "key" and out == "":
        out = norm(values.get("ref"))
    return out


def load_runtime_dict(path: Path) -> Tuple[str, Dict[str, Dict[str, str]]]:
    wb = load_workbook(path)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Missing sheet '{SHEET_NAME}'")
    ws = wb[SHEET_NAME]

    headers: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(row=1, column=c).value)
        if h:
            headers[h] = c

    schema = detect(headers)

    out: Dict[str, Dict[str, str]] = {}
    has_active = "active" in headers
    for r in range(2, ws.max_row + 1):
        values = {name: ws.cell(row=r, column=idx).value for name, idx in headers.items()}
        key = get_val(values, schema, "key")
        en = get_val(values, schema, "en")
        ko = get_val(values, schema, "ko")
        active = parse_active(values.get("active")) if has_active else 1

        if key == "" and en.strip() == "":
            continue
        if active != 1:
            continue

        row = {"en": en}
        if ko:
            row["ko"] = ko
        out[key] = row

    return schema, out


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser(description="Compare effective runtime localization dictionaries between two workbooks.")
    parser.add_argument("--left-xlsx", required=True, help="First workbook path")
    parser.add_argument("--right-xlsx", required=True, help="Second workbook path")
    parser.add_argument("--report-json", default="", help="Optional output report JSON path")
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[2]
    left = (root / args.left_xlsx).resolve() if not Path(args.left_xlsx).is_absolute() else Path(args.left_xlsx)
    right = (root / args.right_xlsx).resolve() if not Path(args.right_xlsx).is_absolute() else Path(args.right_xlsx)

    left_schema, left_dict = load_runtime_dict(left)
    right_schema, right_dict = load_runtime_dict(right)

    left_keys = set(left_dict.keys())
    right_keys = set(right_dict.keys())

    only_left = sorted(left_keys - right_keys)
    only_right = sorted(right_keys - left_keys)

    content_diff = []
    for key in sorted(left_keys & right_keys):
        if left_dict[key] != right_dict[key]:
            content_diff.append(
                {
                    "key": key,
                    "left": left_dict[key],
                    "right": right_dict[key],
                }
            )

    report = {
        "left_schema": left_schema,
        "right_schema": right_schema,
        "left_count": len(left_dict),
        "right_count": len(right_dict),
        "only_left_count": len(only_left),
        "only_right_count": len(only_right),
        "content_diff_count": len(content_diff),
        "only_left_preview": only_left[:20],
        "only_right_preview": only_right[:20],
        "content_diff_preview": content_diff[:20],
    }

    if args.report_json:
        report_path = (root / args.report_json).resolve() if not Path(args.report_json).is_absolute() else Path(args.report_json)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(json.dumps(report, ensure_ascii=False, indent=2))

    if report["only_left_count"] > 0 or report["only_right_count"] > 0 or report["content_diff_count"] > 0:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
