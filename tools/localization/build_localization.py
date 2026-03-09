#!/usr/bin/env python3
"""Validate/build localization JSON from the canonical workbook, with optional Korean autofill."""

from __future__ import annotations

import argparse
import json
import os
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook

SHEET_NAME = "strings"
SCHEMA_LEGACY = "legacy"
SCHEMA_MINIMAL = "minimal"

# Legacy schema (still accepted for compatibility)
LEGACY_REQUIRED_COLUMNS = [
    "key",
    "group",
    "context",
    "source_file",
    "source_line",
    "source_ref",
    "english_en",
    "korean_ko",
    "placeholders",
    "status",
    "notes",
    "active",
    "char_limit",
]

# New minimal schema (canonical)
MINIMAL_REQUIRED_COLUMNS = [
    "ref",
    "en",
    "ko",
]

VALID_STATUS = {"new", "autofilled", "reviewed", "approved"}
PLACEHOLDER_RE = re.compile(r"\{([A-Za-z0-9_]+)\}")

SCHEMA_COLUMN_MAP = {
    SCHEMA_LEGACY: {
        "key": "key",
        "ref": "key",
        "en": "english_en",
        "ko": "korean_ko",
        "status": "status",
        "active": "active",
        "context": "context",
        "char_limit": "char_limit",
        "placeholders": "placeholders",
    },
    SCHEMA_MINIMAL: {
        "key": "key",  # optional; falls back to ref when absent/blank
        "ref": "ref",
        "en": "en",
        "ko": "ko",
        "status": "status",
        "active": "active",
        "context": "context",
        "char_limit": "char_limit",
        "placeholders": "placeholders",
    },
}


@dataclass
class SheetRow:
    row_index: int
    values: Dict[str, object]


def norm_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def text_str(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def parse_active(value: object) -> int:
    s = norm_str(value)
    if s == "":
        return 1
    try:
        if int(float(s)) != 0:
            return 1
        return 0
    except ValueError:
        return -1


def extract_placeholders(text: str) -> List[str]:
    out: List[str] = []
    for m in PLACEHOLDER_RE.finditer(text):
        p = "{" + m.group(1) + "}"
        if p not in out:
            out.append(p)
    return out


def placeholders_csv(text: str) -> str:
    return ",".join(extract_placeholders(text))


def detect_schema(headers: Dict[str, int]) -> str:
    if all(col in headers for col in LEGACY_REQUIRED_COLUMNS):
        return SCHEMA_LEGACY
    if all(col in headers for col in MINIMAL_REQUIRED_COLUMNS):
        return SCHEMA_MINIMAL

    missing_legacy = [c for c in LEGACY_REQUIRED_COLUMNS if c not in headers]
    missing_minimal = [c for c in MINIMAL_REQUIRED_COLUMNS if c not in headers]
    raise ValueError(
        "Workbook schema not recognized.\n"
        f"- Missing legacy columns: {', '.join(missing_legacy)}\n"
        f"- Missing minimal columns: {', '.join(missing_minimal)}"
    )


def column_name(schema: str, logical_name: str) -> str:
    mapping = SCHEMA_COLUMN_MAP.get(schema, {})
    return mapping.get(logical_name, logical_name)


def has_column(headers: Dict[str, int], schema: str, logical_name: str) -> bool:
    col = column_name(schema, logical_name)
    return col in headers


def row_get(row: SheetRow, headers: Dict[str, int], schema: str, logical_name: str) -> str:
    col = column_name(schema, logical_name)
    value = row.values.get(col)
    if logical_name in {"en", "ko", "context"}:
        out = text_str(value)
    else:
        out = norm_str(value)
    if logical_name == "key" and schema == SCHEMA_MINIMAL and out == "":
        out = row_get(row, headers, schema, "ref")
    return out


def row_set(row: SheetRow, headers: Dict[str, int], schema: str, logical_name: str, value: object) -> None:
    col = column_name(schema, logical_name)
    if col in headers:
        row.values[col] = value


def load_sheet_rows(xlsx_path: Path) -> Tuple[object, object, Dict[str, int], List[SheetRow], str]:
    wb = load_workbook(xlsx_path)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Missing sheet '{SHEET_NAME}'")

    ws = wb[SHEET_NAME]
    headers: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = norm_str(ws.cell(row=1, column=c).value)
        if h:
            headers[h] = c

    schema = detect_schema(headers)

    rows: List[SheetRow] = []
    for r in range(2, ws.max_row + 1):
        values = {name: ws.cell(row=r, column=col_idx).value for name, col_idx in headers.items()}
        row = SheetRow(row_index=r, values=values)
        key = row_get(row, headers, schema, "key")
        english = row_get(row, headers, schema, "en")
        if key == "" and english.strip() == "":
            continue
        rows.append(row)

    return wb, ws, headers, rows, schema


def validate_rows(rows: List[SheetRow], headers: Dict[str, int], schema: str) -> List[str]:
    errors: List[str] = []
    seen_keys: Dict[str, int] = {}

    has_status = has_column(headers, schema, "status")
    has_active = has_column(headers, schema, "active")
    has_char_limit = has_column(headers, schema, "char_limit")
    has_placeholders = has_column(headers, schema, "placeholders")

    for row in rows:
        key = row_get(row, headers, schema, "key")
        en = row_get(row, headers, schema, "en")
        ko = row_get(row, headers, schema, "ko")
        status = row_get(row, headers, schema, "status").lower() or "new"
        active = parse_active(row_get(row, headers, schema, "active")) if has_active else 1
        char_limit = row_get(row, headers, schema, "char_limit")
        declared_placeholders = row_get(row, headers, schema, "placeholders")

        if key == "":
            errors.append(f"Row {row.row_index}: key/ref is required")
        elif key in seen_keys:
            errors.append(f"Row {row.row_index}: duplicate key '{key}' (first at row {seen_keys[key]})")
        else:
            seen_keys[key] = row.row_index

        if en.strip() == "":
            errors.append(f"Row {row.row_index}: en text is required")

        if has_status and status not in VALID_STATUS:
            errors.append(
                f"Row {row.row_index}: invalid status '{status}' (allowed: {', '.join(sorted(VALID_STATUS))})"
            )

        if has_active and active not in (0, 1):
            errors.append(f"Row {row.row_index}: active must be 0 or 1")

        en_ph = extract_placeholders(en)
        ko_ph = extract_placeholders(ko)
        if ko and en_ph != ko_ph:
            errors.append(
                f"Row {row.row_index}: placeholder mismatch english={en_ph} korean={ko_ph}"
            )

        if has_placeholders and declared_placeholders:
            canonical = placeholders_csv(en)
            if canonical != declared_placeholders:
                errors.append(
                    f"Row {row.row_index}: placeholders should be '{canonical}' but got '{declared_placeholders}'"
                )

        if has_char_limit and char_limit.strip() != "":
            try:
                max_chars = int(float(char_limit))
                if max_chars > 0 and ko and len(ko) > max_chars:
                    errors.append(
                        f"Row {row.row_index}: ko length {len(ko)} exceeds char_limit {max_chars}"
                    )
            except ValueError:
                errors.append(f"Row {row.row_index}: char_limit must be a number or blank")

    return errors


def _response_text(resp: object) -> str:
    text = getattr(resp, "output_text", None)
    if isinstance(text, str) and text.strip():
        return text.strip()

    output = getattr(resp, "output", None)
    if isinstance(output, list):
        chunks: List[str] = []
        for item in output:
            content = getattr(item, "content", None)
            if not isinstance(content, list):
                continue
            for c in content:
                ctext = getattr(c, "text", None)
                if isinstance(ctext, str):
                    chunks.append(ctext)
        if chunks:
            return "\n".join(chunks).strip()

    raise RuntimeError("Could not parse model response text")


def translate_to_korean(client: object, model: str, english: str, placeholders: List[str], context: str) -> str:
    placeholder_text = ", ".join(placeholders) if placeholders else "(none)"
    system_msg = (
        "Translate game UI/dialogue text to natural Korean. "
        "Preserve placeholders exactly (for example {dmg}, {item}, {qty}). "
        "Return only the Korean translation with no commentary."
    )
    user_msg = (
        f"Context: {context}\n"
        f"English: {english}\n"
        f"Placeholders: {placeholder_text}\n"
        "Korean:"
    )

    if hasattr(client, "responses"):
        resp = client.responses.create(
            model=model,
            input=[
                {"role": "system", "content": system_msg},
                {"role": "user", "content": user_msg},
            ],
            temperature=0,
        )
        return _response_text(resp)

    # compatibility fallback for older SDKs
    resp = client.chat.completions.create(
        model=model,
        temperature=0,
        messages=[
            {"role": "system", "content": system_msg},
            {"role": "user", "content": user_msg},
        ],
    )
    return resp.choices[0].message.content.strip()


def autofill_korean(rows: List[SheetRow], headers: Dict[str, int], schema: str, model: str) -> int:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is required for --autofill-ko")

    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("openai package is required for --autofill-ko") from exc

    client = OpenAI(api_key=api_key)

    updated = 0
    has_active = has_column(headers, schema, "active")
    has_status = has_column(headers, schema, "status")

    for row in rows:
        active = parse_active(row_get(row, headers, schema, "active")) if has_active else 1
        if active != 1:
            continue

        ko = row_get(row, headers, schema, "ko")
        if ko.strip() != "":
            continue

        en = row_get(row, headers, schema, "en")
        if en.strip() == "":
            continue

        placeholders = extract_placeholders(en)
        context = row_get(row, headers, schema, "context")
        if context == "":
            context = row_get(row, headers, schema, "ref")
        translated = translate_to_korean(client, model, en, placeholders, context)
        translated = translated.strip()

        # strict placeholder parity guard
        if extract_placeholders(translated) != placeholders:
            raise RuntimeError(
                f"Autofill placeholder mismatch at row {row.row_index}:\n"
                f"English placeholders={placeholders}\n"
                f"Korean placeholders={extract_placeholders(translated)}"
            )

        row_set(row, headers, schema, "ko", translated)
        status = row_get(row, headers, schema, "status").lower() or "new"
        if has_status and status in {"", "new"}:
            row_set(row, headers, schema, "status", "autofilled")
        updated += 1

    return updated


def write_rows_back(ws: object, headers: Dict[str, int], rows: List[SheetRow]) -> None:
    for row in rows:
        for col_name, col_idx in headers.items():
            if col_name not in row.values:
                continue
            ws.cell(row=row.row_index, column=col_idx).value = row.values[col_name]


def build_json_payload(rows: List[SheetRow], headers: Dict[str, int], schema: str, xlsx_path: Path) -> Dict[str, object]:
    strings: Dict[str, Dict[str, str]] = {}
    has_active = has_column(headers, schema, "active")

    for row in rows:
        active = parse_active(row_get(row, headers, schema, "active")) if has_active else 1
        if active != 1:
            continue

        key = row_get(row, headers, schema, "key")
        en = row_get(row, headers, schema, "en")
        ko = row_get(row, headers, schema, "ko")
        if key == "":
            continue

        entry: Dict[str, str] = {"en": en}
        if ko.strip() != "":
            entry["ko"] = ko
        strings[key] = entry

    ordered = {k: strings[k] for k in sorted(strings.keys())}
    workbook_mtime = datetime.fromtimestamp(xlsx_path.stat().st_mtime, timezone.utc)
    workbook_stamp = workbook_mtime.replace(microsecond=0).isoformat().replace("+00:00", "Z")

    return {
        "meta": {
            "version": 1,
            "source_sheet": SHEET_NAME,
            "source_schema": schema,
            "generated_at_utc": workbook_stamp,
            "row_count": len(ordered),
        },
        "default_lang": "en",
        "languages": ["en", "ko"],
        "strings": ordered,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Build localization JSON from workbook.")
    parser.add_argument("--xlsx", required=True, help="Input workbook path")
    parser.add_argument("--json", required=True, help="Output localization JSON path")
    parser.add_argument("--autofill-ko", action="store_true", help="Autofill blank ko entries using OpenAI API")
    parser.add_argument("--model", default=os.getenv("LOCALIZATION_MODEL", "gpt-4.1-mini"), help="OpenAI model for --autofill-ko")
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[2]
    xlsx_path = (root / args.xlsx).resolve() if not Path(args.xlsx).is_absolute() else Path(args.xlsx)
    json_path = (root / args.json).resolve() if not Path(args.json).is_absolute() else Path(args.json)

    wb, ws, headers, rows, schema = load_sheet_rows(xlsx_path)
    print(f"Detected workbook schema: {schema}")

    if args.autofill_ko:
        updated = autofill_korean(rows, headers, schema, args.model)
        write_rows_back(ws, headers, rows)
        wb.save(xlsx_path)
        print(f"Autofilled ko for {updated} row(s) using model {args.model}")

    errors = validate_rows(rows, headers, schema)
    if errors:
        print("Validation failed:")
        for e in errors:
            print(f"- {e}")
        raise SystemExit(1)

    payload = build_json_payload(rows, headers, schema, xlsx_path)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"Validation succeeded for {len(rows)} row(s)")
    print(f"Wrote localization JSON: {json_path}")


if __name__ == "__main__":
    main()
