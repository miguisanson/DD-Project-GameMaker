#!/usr/bin/env python3
"""Autofill Korean cells in the localization workbook using Google Translate.

This is a draft-pass helper for localization; human review is still required.
"""

from __future__ import annotations

import argparse
import re
import time
from pathlib import Path
from typing import Dict, List, Tuple

from deep_translator import GoogleTranslator
from openpyxl import load_workbook

SHEET_NAME = "strings"
SCHEMA_LEGACY = "legacy"
SCHEMA_MINIMAL = "minimal"
LEGACY_REQUIRED_COLUMNS = ("key", "english_en", "korean_ko")
MINIMAL_REQUIRED_COLUMNS = ("ref", "en", "ko")
PLACEHOLDER_RE = re.compile(r"\{([A-Za-z0-9_]+)\}")

SCHEMA_COLUMN_MAP = {
    SCHEMA_LEGACY: {
        "key": "key",
        "ref": "key",
        "en": "english_en",
        "ko": "korean_ko",
        "status": "status",
        "active": "active",
        "context": "context",
    },
    SCHEMA_MINIMAL: {
        "key": "key",  # optional
        "ref": "ref",
        "en": "en",
        "ko": "ko",
        "status": "status",
        "active": "active",
        "context": "context",
    },
}


def norm_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def text_str(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def detect_schema(headers: Dict[str, int]) -> str:
    if all(col in headers for col in LEGACY_REQUIRED_COLUMNS):
        return SCHEMA_LEGACY
    if all(col in headers for col in MINIMAL_REQUIRED_COLUMNS):
        return SCHEMA_MINIMAL
    raise SystemExit(
        "Workbook schema not recognized. Expected either legacy columns "
        f"{LEGACY_REQUIRED_COLUMNS} or minimal columns {MINIMAL_REQUIRED_COLUMNS}."
    )


def col_name(schema: str, logical_name: str) -> str:
    return SCHEMA_COLUMN_MAP[schema].get(logical_name, logical_name)


def has_col(headers: Dict[str, int], schema: str, logical_name: str) -> bool:
    return col_name(schema, logical_name) in headers


def get_row_value(ws: object, headers: Dict[str, int], schema: str, row: int, logical_name: str) -> str:
    col = col_name(schema, logical_name)
    out = ""
    if col in headers:
        cell_value = ws.cell(row, headers[col]).value
        if logical_name in {"en", "ko", "context"}:
            out = text_str(cell_value)
        else:
            out = norm_str(cell_value)
    if logical_name == "key" and schema == SCHEMA_MINIMAL and out == "":
        out = get_row_value(ws, headers, schema, row, "ref")
    return out


def set_row_value(ws: object, headers: Dict[str, int], schema: str, row: int, logical_name: str, value: object) -> None:
    col = col_name(schema, logical_name)
    if col in headers:
        ws.cell(row, headers[col]).value = value


def parse_active(value: object) -> int:
    s = norm_str(value)
    if s == "":
        return 1
    try:
        return 1 if int(float(s)) != 0 else 0
    except ValueError:
        return 0


def extract_placeholders(text: str) -> List[str]:
    out: List[str] = []
    for m in PLACEHOLDER_RE.finditer(text):
        p = "{" + m.group(1) + "}"
        if p not in out:
            out.append(p)
    return out


def protect_placeholders(text: str) -> Tuple[str, Dict[str, str]]:
    placeholders = extract_placeholders(text)
    protected = text
    token_map: Dict[str, str] = {}
    for i, ph in enumerate(placeholders):
        token = f"ZZPH{i}ZZ"
        token_map[token] = ph
        protected = protected.replace(ph, token)
    return protected, token_map


def restore_placeholders(text: str, token_map: Dict[str, str]) -> str:
    out = text
    for token, ph in token_map.items():
        out = out.replace(token, ph)
    return out


def build_row_text_for_translation(english: str, context: str) -> str:
    # Keep content-only translation to minimize artifacts.
    _ = context
    return english


def main() -> None:
    parser = argparse.ArgumentParser(description="Autofill Korean in game_text.xlsx using Google Translate.")
    parser.add_argument("--xlsx", required=True, help="Workbook path (for example datafiles/localization/game_text.xlsx)")
    parser.add_argument("--sleep-ms", type=int, default=120, help="Delay between requests to reduce throttling risk")
    parser.add_argument("--overwrite", action="store_true", help="Also overwrite non-empty ko values")
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[2]
    xlsx_path = (root / args.xlsx).resolve() if not Path(args.xlsx).is_absolute() else Path(args.xlsx).resolve()
    wb = load_workbook(xlsx_path)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Missing sheet '{SHEET_NAME}' in {xlsx_path}")
    ws = wb[SHEET_NAME]

    headers: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = norm_str(ws.cell(row=1, column=c).value)
        if h:
            headers[h] = c

    schema = detect_schema(headers)
    print(f"Detected workbook schema: {schema}")

    tr = GoogleTranslator(source="en", target="ko")

    updated = 0
    skipped = 0
    failed = 0

    has_active = has_col(headers, schema, "active")
    has_status = has_col(headers, schema, "status")

    for r in range(2, ws.max_row + 1):
        key = get_row_value(ws, headers, schema, r, "key")
        en = get_row_value(ws, headers, schema, r, "en")
        ko = get_row_value(ws, headers, schema, r, "ko")
        context = get_row_value(ws, headers, schema, r, "context") or get_row_value(ws, headers, schema, r, "ref")
        status = get_row_value(ws, headers, schema, r, "status").lower() or "new"
        active = parse_active(get_row_value(ws, headers, schema, r, "active")) if has_active else 1

        if key == "" and en.strip() == "":
            continue
        if active != 1:
            skipped += 1
            continue
        if en.strip() == "":
            skipped += 1
            continue
        if ko.strip() != "" and not args.overwrite:
            skipped += 1
            continue

        source = build_row_text_for_translation(en, context)
        protected, token_map = protect_placeholders(source)

        try:
            translated = tr.translate(protected)
            if translated is None:
                translated = ""
            translated = restore_placeholders(norm_str(translated), token_map)

            # Hard guard on placeholder parity.
            if extract_placeholders(translated) != extract_placeholders(en):
                failed += 1
                continue

            set_row_value(ws, headers, schema, r, "ko", translated)
            if has_status and status in ("", "new"):
                set_row_value(ws, headers, schema, r, "status", "autofilled")
            updated += 1
        except Exception:
            failed += 1

        if args.sleep_ms > 0:
            time.sleep(args.sleep_ms / 1000.0)

    wb.save(xlsx_path)
    print(f"Autofill finished for {xlsx_path}")
    print(f"updated={updated} skipped={skipped} failed={failed}")


if __name__ == "__main__":
    main()
