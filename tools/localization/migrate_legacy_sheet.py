#!/usr/bin/env python3
"""Migrate legacy localization workbook schema to minimal ref/en/ko schema.

Legacy expected columns:
- key, group, context, source_file, source_line, source_ref,
  english_en, korean_ko, placeholders, status, notes, active, char_limit

New minimal columns (default):
- ref, en, ko
Optional:
- key (via --include-key)
- notes (via --include-notes)
"""

from __future__ import annotations

import argparse
import csv
import json
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook, load_workbook

SHEET_NAME = "strings"
LEGACY_REQUIRED_COLUMNS = [
    "key",
    "group",
    "context",
    "source_file",
    "source_line",
    "source_ref",
    "english_en",
    "korean_ko",
    "placeholders",
    "status",
    "notes",
    "active",
    "char_limit",
]


@dataclass
class LegacyRow:
    row_index: int
    values: Dict[str, object]


def norm_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def text_str(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def extract_placeholders(text: str) -> List[str]:
    out: List[str] = []
    i = 0
    while i < len(text):
        if text[i] == "{":
            end = text.find("}", i + 1)
            if end != -1:
                token = text[i : end + 1]
                # Keep strict token shape for parity checks.
                if token.startswith("{") and token.endswith("}") and len(token) > 2 and token[1:-1].replace("_", "a").isalnum():
                    if token not in out:
                        out.append(token)
                i = end + 1
                continue
        i += 1
    return out


def check_balanced_braces(text: str) -> bool:
    depth = 0
    for ch in text:
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth < 0:
                return False
    return depth == 0


def load_legacy_rows(path: Path) -> Tuple[object, object, Dict[str, int], List[LegacyRow]]:
    wb = load_workbook(path)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Missing sheet '{SHEET_NAME}'")

    ws = wb[SHEET_NAME]
    headers: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = norm_str(ws.cell(row=1, column=c).value)
        if h:
            headers[h] = c

    missing = [c for c in LEGACY_REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise ValueError(
            "Input workbook does not match legacy schema. "
            f"Missing required columns: {', '.join(missing)}"
        )

    rows: List[LegacyRow] = []
    for r in range(2, ws.max_row + 1):
        values = {name: ws.cell(row=r, column=col_idx).value for name, col_idx in headers.items()}
        if norm_str(values.get("key")) == "" and text_str(values.get("english_en")).strip() == "":
            continue
        rows.append(LegacyRow(row_index=r, values=values))

    return wb, ws, headers, rows


def choose_unique_ref(base_ref: str, seen: Dict[str, int]) -> str:
    ref = base_ref
    if ref == "":
        ref = "missing_ref"

    if ref not in seen:
        seen[ref] = 1
        return ref

    seen[ref] += 1
    return f"{ref}#{seen[ref]}"


def migrate_rows(
    rows: List[LegacyRow],
    include_key: bool,
    include_notes: bool,
) -> Tuple[List[Dict[str, str]], List[Dict[str, str]], Dict[str, object]]:
    migrated_rows: List[Dict[str, str]] = []
    mapping_rows: List[Dict[str, str]] = []

    seen_refs: Dict[str, int] = {}
    duplicate_refs: List[Dict[str, str]] = []
    missing_en: List[Dict[str, str]] = []
    missing_ko: List[Dict[str, str]] = []
    placeholder_mismatch: List[Dict[str, str]] = []
    markup_issues: List[Dict[str, str]] = []

    for row in rows:
        legacy_key = norm_str(row.values.get("key"))
        en = text_str(row.values.get("english_en"))
        ko = text_str(row.values.get("korean_ko"))
        context = text_str(row.values.get("context"))
        source_file = text_str(row.values.get("source_file"))
        source_line = text_str(row.values.get("source_line"))
        source_ref = text_str(row.values.get("source_ref"))

        new_ref = choose_unique_ref(legacy_key, seen_refs)
        if new_ref != legacy_key:
            duplicate_refs.append(
                {
                    "legacy_row": str(row.row_index),
                    "legacy_key": legacy_key,
                    "resolved_ref": new_ref,
                }
            )

        out_row: Dict[str, str] = {
            "ref": new_ref,
            "en": en,
            "ko": ko,
        }
        if include_key:
            out_row["key"] = legacy_key or new_ref
        if include_notes:
            out_row["notes"] = f"{context} | {source_file}:{source_line} | {source_ref}".strip(" |")
        migrated_rows.append(out_row)

        mapping_rows.append(
            {
                "legacy_row": str(row.row_index),
                "legacy_key": legacy_key,
                "new_ref": new_ref,
                "runtime_key": (legacy_key or new_ref),
                "group": text_str(row.values.get("group")),
                "context": context,
                "source_file": source_file,
                "source_line": source_line,
                "source_ref": source_ref,
            }
        )

        if en.strip() == "":
            missing_en.append(
                {
                    "legacy_row": str(row.row_index),
                    "legacy_key": legacy_key,
                    "new_ref": new_ref,
                }
            )
        if ko.strip() == "":
            missing_ko.append(
                {
                    "legacy_row": str(row.row_index),
                    "legacy_key": legacy_key,
                    "new_ref": new_ref,
                }
            )

        en_ph = extract_placeholders(en)
        ko_ph = extract_placeholders(ko)
        if ko and en_ph != ko_ph:
            placeholder_mismatch.append(
                {
                    "legacy_row": str(row.row_index),
                    "legacy_key": legacy_key,
                    "new_ref": new_ref,
                    "en_placeholders": ",".join(en_ph),
                    "ko_placeholders": ",".join(ko_ph),
                }
            )

        if not check_balanced_braces(en) or (ko and not check_balanced_braces(ko)):
            markup_issues.append(
                {
                    "legacy_row": str(row.row_index),
                    "legacy_key": legacy_key,
                    "new_ref": new_ref,
                    "issue": "unbalanced_braces",
                }
            )

    report = {
        "input_row_count": len(rows),
        "output_row_count": len(migrated_rows),
        "duplicate_ref_resolutions": duplicate_refs,
        "missing_en": missing_en,
        "missing_ko": missing_ko,
        "placeholder_mismatch": placeholder_mismatch,
        "markup_issues": markup_issues,
        "summary": {
            "duplicate_ref_resolutions": len(duplicate_refs),
            "missing_en": len(missing_en),
            "missing_ko": len(missing_ko),
            "placeholder_mismatch": len(placeholder_mismatch),
            "markup_issues": len(markup_issues),
        },
    }
    return migrated_rows, mapping_rows, report


def write_minimal_workbook(path: Path, rows: List[Dict[str, str]], include_key: bool, include_notes: bool) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    columns = ["ref", "en", "ko"]
    if include_key:
        columns.append("key")
    if include_notes:
        columns.append("notes")

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(columns)

    for row in rows:
        ws.append([row.get(col, "") for col in columns])

    ws.freeze_panes = "A2"
    last_col_letter = chr(ord("A") + len(columns) - 1)
    ws.auto_filter.ref = f"A1:{last_col_letter}{ws.max_row}"

    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 72
    if include_key:
        ws.column_dimensions["D"].width = 52
    if include_notes:
        ws.column_dimensions["E" if include_key else "D"].width = 64

    wb.save(path)


def write_mapping_csv(path: Path, mapping_rows: List[Dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = [
        "legacy_row",
        "legacy_key",
        "new_ref",
        "runtime_key",
        "group",
        "context",
        "source_file",
        "source_line",
        "source_ref",
    ]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        for row in mapping_rows:
            writer.writerow({k: row.get(k, "") for k in columns})


def write_report(path: Path, report: Dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Migrate legacy localization workbook to minimal schema.")
    parser.add_argument("--input-xlsx", required=True, help="Legacy workbook path")
    parser.add_argument("--output-xlsx", required=True, help="Output workbook path (minimal schema)")
    parser.add_argument("--mapping-csv", required=True, help="Output mapping CSV path")
    parser.add_argument("--report-json", required=True, help="Output validation report JSON path")
    parser.add_argument("--backup-xlsx", default="", help="Optional backup copy path for input workbook")
    parser.add_argument("--include-key", action="store_true", help="Include optional key column in output workbook")
    parser.add_argument("--include-notes", action="store_true", help="Include optional notes column in output workbook")
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[2]

    input_xlsx = (root / args.input_xlsx).resolve() if not Path(args.input_xlsx).is_absolute() else Path(args.input_xlsx)
    output_xlsx = (root / args.output_xlsx).resolve() if not Path(args.output_xlsx).is_absolute() else Path(args.output_xlsx)
    mapping_csv = (root / args.mapping_csv).resolve() if not Path(args.mapping_csv).is_absolute() else Path(args.mapping_csv)
    report_json = (root / args.report_json).resolve() if not Path(args.report_json).is_absolute() else Path(args.report_json)

    if args.backup_xlsx:
        backup_xlsx = (root / args.backup_xlsx).resolve() if not Path(args.backup_xlsx).is_absolute() else Path(args.backup_xlsx)
        backup_xlsx.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(input_xlsx, backup_xlsx)
        print(f"Wrote backup workbook: {backup_xlsx}")

    _, _, _, rows = load_legacy_rows(input_xlsx)
    migrated_rows, mapping_rows, report = migrate_rows(rows, include_key=args.include_key, include_notes=args.include_notes)

    write_minimal_workbook(output_xlsx, migrated_rows, include_key=args.include_key, include_notes=args.include_notes)
    write_mapping_csv(mapping_csv, mapping_rows)
    write_report(report_json, report)

    print(f"Migrated {len(rows)} row(s) from legacy schema -> minimal schema")
    print(f"Output workbook: {output_xlsx}")
    print(f"Mapping CSV: {mapping_csv}")
    print(f"Validation report: {report_json}")
    print(f"Summary: {report['summary']}")


if __name__ == "__main__":
    main()
