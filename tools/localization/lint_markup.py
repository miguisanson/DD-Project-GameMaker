#!/usr/bin/env python3
"""Check the text markup (D149, D156) in every workbook string. Exit 1 when something is wrong.

The rules are in docs/text-markup.md; the game's parser is Markup_Parse in scripts/scr_input/scr_input.gml and this file follows it.
Flagged, per string and language:
  - a tag the game drops: unknown (`[x]`, `[k:dance]`), unclosed (`[t]` with no `[/t]`), nested (a tag inside a tag), stray (`[/t]` with no opener);
  - the old `{confirm}` spelling of a button (write `[k:confirm]`);
  - [a] around anything but BREAK, Perfect (or their Korean words) (the one gold accent is for rare / critical words);
  - an item tag ([r1] [r2] [r3], D178: an item's name in its rarity colour) written in the workbook: the game puts it around every item name it fills into
    {item} / {name} (Item_NameMarked), so a string never carries one;
  - English and Korean that do not carry the same number of [a] spans and [k:..] buttons;
  - key letters written as text in English (A/D, WASD, Press Z, Up / Down ...): use [k:left][k:right], [k:confirm].

Usage:
  python tools/localization/lint_markup.py [--xlsx datafiles/localization/game_text.xlsx]
  python tools/localization/lint_markup.py --selftest
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import List, Tuple

SHEET = "strings"
ACTIONS = {"confirm", "back", "cancel", "up", "down", "left", "right", "interact", "run", "tab_prev", "tab_next", "menu", "pause", "start"}
ACCENT_OK = {"BREAK", "Perfect", "브레이크", "완벽"}
ITEM_TAGS = ("r1", "r2", "r3")
KEY_LETTERS = ["A/D", "W/S", "WASD", "A / D", "Up / Down", "Left / Right", "Press Z", "Press X", "Press Space", "Press Enter", "Press Esc"]
OLD_BUTTON_RE = re.compile(r"\{(confirm|back|cancel|up|down|left|right|interact|run|tab_prev|tab_next|menu|pause|start)\}")


def tag_kind(inner: str) -> str:
    """What the text between [ and ] is: t / a / /t / /a / k:<action> / ? (tag-shaped, unknown) / '' (not a tag)."""
    m = re.fullmatch(r"(/?)([a-z]+|r[0-9]+)(?::([a-z_]+))?", inner)
    if not m:
        return ""
    close, name, arg = m.group(1), m.group(2), m.group(3)
    if close:
        if arg is not None:
            return ""
        return "/" + name if name in ("t", "a") + ITEM_TAGS else "?"
    if arg is not None:
        return "k:" + arg if name == "k" and arg in ACTIONS else "?"
    return name if name in ("t", "a") + ITEM_TAGS else "?"


def has_tag(text: str) -> bool:
    pos = 0
    while True:
        o = text.find("[", pos)
        if o < 0:
            return False
        c = text.find("]", o + 1)
        if c > o + 1 and tag_kind(text[o + 1:c]) != "":
            return True
        pos = o + 1


def parse(text: str) -> Tuple[List[Tuple[str, str]], List[str]]:
    """Returns ([(kind, text)], problems): kind is 'text', 'term', 'accent', 'item' or 'action'. Mirrors Markup_Parse."""
    out: List[Tuple[str, str]] = []
    problems: List[str] = []
    buf = ""
    s = text
    while s:
        o = s.find("[")
        if o < 0:
            buf += s
            break
        rest = s[o + 1:]
        c = rest.find("]")
        kind = tag_kind(rest[:c]) if c > 0 else ""
        if kind == "":
            buf += s[:o + 1]
            s = rest
            continue
        buf += s[:o]
        after = rest[c + 1:]
        if kind in ("t", "a") + ITEM_TAGS:
            cp = after.find("[/" + kind + "]")
            body = after[:cp] if cp >= 0 else ""
            if cp >= 0 and not has_tag(body):
                if buf:
                    out.append(("text", buf))
                buf = ""
                if body:
                    out.append(("term" if kind == "t" else ("accent" if kind == "a" else "item"), body))
                s = after[cp + 3 + len(kind):]
                continue
            problems.append(("nested [%s]" if cp >= 0 else "unclosed [%s]") % kind)
        elif kind.startswith("k:"):
            if buf:
                out.append(("text", buf))
            buf = ""
            out.append(("action", kind[2:]))
            s = after
            continue
        elif kind == "?":
            problems.append("unknown tag [" + rest[:c] + "]")
        else:
            problems.append("stray [" + kind + "]")
        s = after
    if buf:
        out.append(("text", buf))
    return out, problems


def check_row(ref: str, en: str, ko: str) -> List[str]:
    errs: List[str] = []
    counts = {}
    for lang, text in (("en", en), ("ko", ko)):
        toks, probs = parse(text)
        for p in probs:
            errs.append("%s [%s]: %s" % (ref, lang, p))
        if OLD_BUTTON_RE.search(text):
            errs.append("%s [%s]: old {action} button, write [k:action]" % (ref, lang))
        for k, t in toks:
            if k == "item":
                errs.append("%s [%s]: item tag around '%s' (the game marks item names itself; write plain {item})" % (ref, lang, t))
        accents = [t for k, t in toks if k == "accent"]
        for a in accents:
            if a not in ACCENT_OK:
                errs.append("%s [%s]: [a] around '%s' (only BREAK, Perfect)" % (ref, lang, a))
        counts[lang] = (len(accents), sum(1 for k, _ in toks if k == "action"))
    if counts["en"] != counts["ko"]:
        errs.append("%s: en has %d accents / %d buttons, ko has %d / %d" % (ref, counts["en"][0], counts["en"][1], counts["ko"][0], counts["ko"][1]))
    for pat in KEY_LETTERS:
        if pat in en:
            errs.append("%s [en]: key letters in text ('%s'), use [k:..]" % (ref, pat))
    return errs


def selftest() -> int:
    bad = 0

    def expect(name: str, got, want) -> None:
        nonlocal bad
        if got != want:
            bad += 1
            print("SELFTEST FAIL %s: got %r want %r" % (name, got, want))

    toks, probs = parse("[t]Poison[/t] hurts")
    expect("term then plain", toks, [("term", "Poison"), ("text", " hurts")])
    expect("clean", probs, [])
    toks, probs = parse("[t]a [a]b[/a] c[/t]")
    expect("nested drops the outer tag", [k for k, _ in toks], ["text", "accent", "text"])
    expect("nested is a problem", len(probs) >= 1, True)
    expect("unknown tag", len(parse("a [zz]b[/zz]")[1]) >= 1, True)
    expect("unclosed", parse("[t]x")[1], ["unclosed [t]"])
    expect("unknown action", len(parse("[k:dance]")[1]), 1)
    expect("a lone bracket is text", parse("a [ b ] c"), ([("text", "a [ b ] c")], []))
    expect("button", parse("Press [k:confirm].")[0], [("text", "Press "), ("action", "confirm"), ("text", ".")])
    expect("old button flagged", len(check_row("x", "Press {confirm}", "{confirm}")) >= 1, True)
    expect("accent allow-list", len(check_row("x", "a [a]gold[/a]", "a [a]gold[/a]")) >= 1, True)
    expect("parity", len(check_row("x", "[a]BREAK[/a]", "BREAK")) >= 1, True)
    expect("key letters", len(check_row("x", "Move with A/D", "A/D")) >= 1, True)
    expect("item tags parse", parse("Got [r2]Iron Sword[/r2] x1")[0], [("text", "Got "), ("item", "Iron Sword"), ("text", " x1")])
    expect("item tags: clean", parse("[r1]a[/r1] [r3]b[/r3]")[1], [])
    expect("[r4] is unknown", len(parse("[r4]a[/r4]")[1]) >= 1, True)
    expect("a tag in an item tag nests", len(parse("[r2]a [t]b[/t][/r2]")[1]) >= 1, True)
    expect("an item tag in the workbook is flagged", len(check_row("x", "Got [r2]{item}[/r2]", "[r2]{item}[/r2]")) >= 1, True)
    expect("plain {item} is fine", check_row("x", "Got {item}", "{item}"), [])
    expect("[a] around {item} is flagged now", len(check_row("x", "Got [a]{item}[/a]", "[a]{item}[/a]")) >= 1, True)
    expect("a good row", check_row("x", "[t]Poison[/t] and [a]BREAK[/a]: [k:confirm]", "[t]독[/t]과 [a]BREAK[/a]: [k:confirm]"), [])
    print("selftest %s" % ("FAILED" if bad else "ok"))
    return 1 if bad else 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default="datafiles/localization/game_text.xlsx")
    ap.add_argument("--selftest", action="store_true")
    args = ap.parse_args()
    if args.selftest:
        return selftest()
    from openpyxl import load_workbook

    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    ws = wb[SHEET]
    errs: List[str] = []
    rows = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not row or not row[0]:
            continue
        ref = str(row[0]).strip()
        en = "" if row[1] is None else str(row[1])
        ko = "" if row[2] is None else str(row[2])
        rows += 1
        errs.extend(check_row(ref, en, ko))
    for e in errs:
        print("MARKUP " + e)
    print("markup-lint: %d row(s), %d problem(s)" % (rows, len(errs)))
    return 1 if errs else 0


if __name__ == "__main__":
    sys.exit(main())
